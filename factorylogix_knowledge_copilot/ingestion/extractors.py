"""Extracción segura de texto por tipo de archivo, con límites de páginas/filas y timeout."""
from __future__ import annotations

import concurrent.futures
import csv
import importlib.util
import io
import shutil
from dataclasses import dataclass, field
from html.parser import HTMLParser

from core.errors import ExtractionError


@dataclass
class PageText:
    page: int | None
    text: str


@dataclass
class ExtractedDocument:
    pages: list[PageText]
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf(data: bytes, max_pages: int) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, ValueError, OSError) as exc:
        raise ExtractionError("PDF dañado o ilegible.") from exc
    if reader.is_encrypted:
        raise ExtractionError("PDF cifrado: no se procesan documentos protegidos con contraseña.")
    total = len(reader.pages)
    if total > max_pages:
        raise ExtractionError(f"El PDF tiene {total} páginas; el máximo permitido es {max_pages}.")
    pages, warnings = [], []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            pages.append(PageText(idx, page.extract_text() or ""))
        except Exception:  # noqa: BLE001 - una página corrupta no debe detener todo el documento
            warnings.append(f"No se pudo leer la página {idx}.")
    if not any(p.text.strip() for p in pages):
        warnings.append("El PDF no contiene texto seleccionable (posible escaneo). Usa OCR si está disponible.")
    return ExtractedDocument(pages=pages, page_count=total, warnings=warnings)


def extract_docx(data: bytes) -> ExtractedDocument:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx lanza varios tipos
        raise ExtractionError("DOCX dañado o ilegible.") from exc
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style.startswith("título") or style.startswith("titulo"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = min(int(digits) if digits else 1, 3)
            lines.append("#" * level + " " + text)
        elif style.startswith("title"):
            lines.append("# " + text)
        elif "list" in style:
            lines.append("- " + text)
        else:
            lines.append(text)
    for t_index, table in enumerate(document.tables, start=1):
        lines.append(f"\n### Tabla {t_index}")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" | ".join(cells))
    return ExtractedDocument(pages=[PageText(None, "\n".join(lines))])


def extract_xlsx(data: bytes, max_rows: int) -> ExtractedDocument:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("XLSX dañado o ilegible.") from exc
    pages, warnings, total_rows = [], [], 0
    try:
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(h).strip() if h is not None else f"col{i + 1}" for i, h in enumerate(rows[0])]
            lines = [f"## Hoja: {sheet.title}"]
            for row in rows[1:]:
                total_rows += 1
                if total_rows > max_rows:
                    warnings.append(f"Se truncó la lectura a {max_rows} filas.")
                    break
                pairs = [f"{header[i] if i < len(header) else f'col{i + 1}'}: {v}"
                         for i, v in enumerate(row) if v not in (None, "")]
                if pairs:
                    lines.append("; ".join(pairs))
            pages.append(PageText(None, "\n".join(lines)))
            if total_rows > max_rows:
                break
    finally:
        wb.close()
    return ExtractedDocument(pages=pages, warnings=warnings)


_QUESTION_COLS = {"pregunta", "question", "q"}
_ANSWER_COLS = {"respuesta", "answer", "a"}


def extract_csv(data: bytes, max_rows: int) -> ExtractedDocument:
    text = _decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = []
    warnings = []
    header: list[str] = []
    q_col = a_col = -1
    for idx, row in enumerate(reader):
        if idx == 0:
            header = [h.strip() for h in row]
            lowered = [h.lower() for h in header]
            q_col = next((i for i, h in enumerate(lowered) if h in _QUESTION_COLS), -1)
            a_col = next((i for i, h in enumerate(lowered) if h in _ANSWER_COLS), -1)
            continue
        if idx > max_rows:
            warnings.append(f"Se truncó la lectura a {max_rows} filas.")
            break
        if q_col >= 0 and a_col >= 0 and len(row) > max(q_col, a_col) and row[q_col].strip():
            # Formato FAQ: cada fila se convierte en un tema con su respuesta breve.
            extra = [f"{header[i]}: {v.strip()}" for i, v in enumerate(row)
                     if i not in (q_col, a_col) and i < len(header) and v.strip()]
            rows.append(f"## {row[q_col].strip()}\n### Respuesta breve\n{row[a_col].strip()}"
                        + (f"\n{'; '.join(extra)}" if extra else "") + "\n")
            continue
        pairs = [f"{header[i] if i < len(header) else f'col{i + 1}'}: {v.strip()}" for i, v in enumerate(row)
                 if v.strip()]
        if pairs:
            rows.append("; ".join(pairs))
    return ExtractedDocument(pages=[PageText(None, "\n".join(rows))], warnings=warnings)


class _HTMLToText(HTMLParser):
    _SKIP = {"script", "style", "noscript", "iframe", "object", "embed", "svg", "head"}
    _BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "table", "ul", "ol"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip_depth = 0
        self.heading = ""

    def handle_starttag(self, tag, attrs):  # noqa: ANN001
        if tag in self._SKIP:
            self.skip_depth += 1
        elif tag in {"h1", "h2", "h3"}:
            self.parts.append("\n" + "#" * int(tag[1]) + " ")
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):  # noqa: ANN001
        if tag in self._SKIP and self.skip_depth:
            self.skip_depth -= 1
        elif tag in {"h1", "h2", "h3"} or tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):  # noqa: ANN001
        if not self.skip_depth:
            self.parts.append(data)


def extract_html(data: bytes) -> ExtractedDocument:
    parser = _HTMLToText()
    parser.feed(_decode_text(data))
    parser.close()
    return ExtractedDocument(pages=[PageText(None, "".join(parser.parts))])


def ocr_available() -> bool:
    return importlib.util.find_spec("pytesseract") is not None and shutil.which("tesseract") is not None


def extract_image(data: bytes) -> ExtractedDocument:
    if not ocr_available():
        raise ExtractionError("OCR no disponible: instala Tesseract y pytesseract (opcional) o carga un PDF con texto.")
    import pytesseract
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    text = pytesseract.image_to_string(image, lang="spa+eng")
    return ExtractedDocument(pages=[PageText(1, text)], page_count=1,
                             warnings=["Texto obtenido por OCR: revisar cuidadosamente antes de aprobar."])


def extract(kind: str, data: bytes, *, max_pages: int, max_rows: int, timeout: float) -> ExtractedDocument:
    """Ejecuta la extracción con timeout. Nota: un hilo que excede el tiempo no puede forzarse a terminar
    en Python; se abandona su resultado y se reporta el error al usuario."""
    funcs = {
        "pdf": lambda: extract_pdf(data, max_pages),
        "docx": lambda: extract_docx(data),
        "xlsx": lambda: extract_xlsx(data, max_rows),
        "csv": lambda: extract_csv(data, max_rows),
        "text": lambda: ExtractedDocument(pages=[PageText(None, _decode_text(data))]),
        "markdown": lambda: ExtractedDocument(pages=[PageText(None, _decode_text(data))]),
        "html": lambda: extract_html(data),
        "image": lambda: extract_image(data),
    }
    if kind not in funcs:
        raise ExtractionError("Tipo de archivo no soportado.")
    executor = concurrent.futures.ThreadPoolExecutor(max_workers=1)
    future = executor.submit(funcs[kind])
    try:
        return future.result(timeout=timeout)
    except concurrent.futures.TimeoutError as exc:
        raise ExtractionError(f"La extracción excedió {timeout:g} segundos.") from exc
    finally:
        executor.shutdown(wait=False, cancel_futures=True)
