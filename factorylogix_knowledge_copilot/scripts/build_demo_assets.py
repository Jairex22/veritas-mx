"""Genera los documentos DEMO binarios (DOCX y PDF) a partir de texto incluido en este script.

Uso: python scripts/build_demo_assets.py
Los archivos resultantes están marcados como DEMO y no contienen información real.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "demo_data" / "documents"

MANUAL_SECTIONS = [
    ("Dónde consultar la genealogía", "genealogía, genealogy, assembly genealogy, trazabilidad",
     "La genealogía de un ensamble se consulta en FactoryLogix Analytics o en la vista de trazabilidad, "
     "buscando por el número de serie del ensamble padre. Muestra subensambles y componentes instalados.",
     ["Escanea o escribe el número de serie del ensamble.",
      "Abre la consulta de genealogía o trazabilidad del producto.",
      "Revisa la relación padre/hijo de subensambles.",
      "Revisa los componentes instalados con su número de parte y lote."],
     "Escala a Soporte MES si la genealogía aparece incompleta o con un subensamble que no corresponde."),
    ("Cómo consultar WIP", "wip, work in process, trabajo en proceso, inventario en proceso",
     "El WIP (trabajo en proceso) muestra cuántas unidades hay en cada operación. Se consulta en los reportes de "
     "Analytics filtrando por línea, Work Order y operación.",
     ["Abre los reportes de Analytics.", "Selecciona el reporte de WIP.",
      "Filtra por línea y Work Order.", "Identifica operaciones con acumulación de unidades."],
     "Escala al supervisor si detectas unidades acumuladas o estancadas en una operación."),
    ("Componentes instalados", "componentes instalados, installed components, lote, número de parte",
     "Los componentes instalados registran qué número de parte y lote se colocó en cada unidad, lo que permite "
     "trazabilidad ante problemas de calidad.",
     ["Consulta el serial de la unidad.", "Revisa la lista de componentes instalados.",
      "Compara número de parte y lote contra la BOM."],
     "Escala a Calidad si un lote instalado no corresponde a la BOM vigente."),
]

PDF_PAGES = [
    ["DEMO - FactoryLogix Production quick reference (DEMO, not official)",
     "",
     "Serialization overview",
     "Keywords: serialization, serial number, unique identifier",
     "Serialization assigns a unique serial number to every unit so that each operation,",
     "test result and installed component can be traced across the production route.",
     "Check that the serial label is readable before the first operation.",
     "Escalate to MES Support if duplicate serial numbers are detected."],
    ["Pass and Fail results",
     "Keywords: pass, fail, test result, failed unit",
     "A Pass result means the unit met the test or inspection criteria.",
     "A Fail result means the unit did not meet the criteria and a symptom must be recorded",
     "before the unit is sent to Repair. Never mark a failed unit as Pass to keep production moving.",
     "Escalate to Quality if the same failure repeats on several units."],
]


def build_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("DEMO - Manual: Genealogía, trazabilidad y WIP", level=1)
    document.add_paragraph("DOCUMENTO DEMO. Manual ilustrativo; no es un manual oficial de FactoryLogix.")
    for title, keywords, brief, steps, escalate in MANUAL_SECTIONS:
        document.add_heading(title, level=2)
        document.add_paragraph(f"Palabras clave: {keywords}")
        document.add_heading("Respuesta breve", level=3)
        document.add_paragraph(brief)
        document.add_heading("Pasos", level=3)
        for i, step in enumerate(steps, start=1):
            document.add_paragraph(f"{i}. {step}")
        document.add_heading("Cuándo escalar", level=3)
        document.add_paragraph(escalate)
    document.save(path)


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf(path: Path) -> None:
    """PDF mínimo válido (Helvetica, WinAnsi) escrito a mano para no depender de librerías extra."""
    objects: list[bytes] = []
    page_ids = []
    font_id = 3
    n_pages = len(PDF_PAGES)
    first_page_obj = 4
    for i, lines in enumerate(PDF_PAGES):
        page_obj = first_page_obj + i * 2
        content_obj = page_obj + 1
        page_ids.append(page_obj)
        stream_lines = ["BT", "/F1 11 Tf", "14 TL", "60 780 Td"]
        for line in lines:
            stream_lines.append(f"({_pdf_escape(line)}) Tj T*")
        stream_lines.append("ET")
        stream = "\n".join(stream_lines).encode("cp1252")
        objects.append((page_obj, f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
                                  f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_obj} 0 R >>".encode()))
        objects.append((content_obj, b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream +
                        b"\nendstream"))
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    objects = [(1, b"<< /Type /Catalog /Pages 2 0 R >>"),
               (2, f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode()),
               (3, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")] + objects
    objects.sort(key=lambda o: o[0])
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for num, body in objects:
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for num in range(1, len(objects) + 1):
        out += f"{offsets[num]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(OUT / "DEMO-MAN-005_genealogia_wip.docx")
    build_pdf(OUT / "DEMO-MAN-014_production_quick_reference.pdf")
    print("Documentos DEMO generados en", OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
