#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GPV PRODUCT CONFIGURATION
=========================

Desktop viewer (tkinter + ttk) for the SMT marking validation workbook
"EMX-KA661_Rev.0 Validacion de marcado SMT (1).xlsm".

* The workbook is opened strictly READ-ONLY. The application never writes,
  saves or re-packs the file, so macros (vbaProject.bin) stay untouched.
  Every load works on a temporary snapshot copy, so Excel may keep the
  original open (and engineering may save it) while the viewer runs.
* openpyxl is used when available; otherwise a built-in reader based only on
  zipfile + xml.etree.ElementTree reads cells and sharedStrings.
* Headers are detected automatically (alias system, see RecordNormalizer).
* The BOM tree is built from real Parent/Level columns when the workbook has
  them populated. Otherwise a LOGICAL VIEW (Item -> Alternatives ->
  Manufacturer / MPN / Marking) is shown. The logical view is NOT a
  FactoryLogix genealogy.

Run:  python GPV_Product_Configuration.py
      python GPV_Product_Configuration.py --check      (console data report)
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import logging
import math
import os
import queue
import re
import shutil
import sys
import tempfile
import threading
import time
import zipfile
import xml.etree.ElementTree as ET
from collections import OrderedDict
from logging.handlers import RotatingFileHandler
from pathlib import Path

try:  # The data layer (and --check) works without tkinter.
    import tkinter as tk
    from tkinter import ttk, messagebox, font as tkfont
except ImportError:  # pragma: no cover - depends on the Python install
    tk = None

try:
    import openpyxl  # optional
except Exception:  # pragma: no cover - optional dependency
    openpyxl = None


# ============================================================================
# Constants / paths
# ============================================================================

APP_TITLE = "GPV PRODUCT CONFIGURATION"
APP_VERSION = "1.0.0"
BASE_DIR = Path(__file__).resolve().parent
CONFIG_FILE = BASE_DIR / "config" / "settings.json"
LOG_DIR = BASE_DIR / "logs"
LOG_FILE = LOG_DIR / "app.log"
EMPTY = "—"  # em dash shown for any missing value

DEFAULT_SETTINGS = {
    "excel_file": "data/EMX-KA661_Rev.0 Validacion de marcado SMT (1).xlsm",
    "default_mode": "ENGINEERING",
    "window_width": 850,
    "window_height": 850,
    "reader": "auto",            # auto | openpyxl | builtin
    "max_tree_items": 150,       # max. items listed for a partial search
    "details_max_rows": 5000,    # max. rows rendered in VIEW DETAILS
    "extra_column_aliases": {},  # {"mpn": ["my mpn header"], ...}
}

# Palette (from the reference image / specification)
C_NAVY = "#202830"
C_GREEN = "#68B82E"
C_BLUE = "#2979A8"
C_BLUE_DARK = "#1F6690"
C_BG = "#EEF2F4"
C_PANEL = "#FFFFFF"
C_BORDER = "#AEB7BF"
C_TEXT = "#1A1F24"
C_MUTED = "#66717B"
C_SOFT = "#DDE3E7"
C_RED = "#C62828"
C_AMBER = "#A15C00"
C_OK = "#2E7D32"
C_SELECT = "#D6E8F3"


# ============================================================================
# Settings / logging
# ============================================================================

def load_settings(path: Path = CONFIG_FILE) -> dict:
    settings = dict(DEFAULT_SETTINGS)
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        if isinstance(data, dict):
            settings.update(data)
    except FileNotFoundError:
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(DEFAULT_SETTINGS, fh, indent=2, ensure_ascii=False)
    except (OSError, ValueError) as exc:
        logging.getLogger("gpv").error("settings.json could not be read: %s", exc)
    return settings


def resolve_path(value: str) -> Path:
    """Relative paths in settings.json are relative to the application folder."""
    p = Path(value)
    return p if p.is_absolute() else (BASE_DIR / p)


def setup_logging() -> logging.Logger:
    logger = logging.getLogger("gpv")
    if logger.handlers:
        return logger
    logger.setLevel(logging.INFO)
    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", "%Y-%m-%d %H:%M:%S")
    try:
        LOG_DIR.mkdir(parents=True, exist_ok=True)
        handler = RotatingFileHandler(LOG_FILE, maxBytes=2_000_000, backupCount=3, encoding="utf-8")
        handler.setFormatter(fmt)
        logger.addHandler(handler)
    except OSError:
        logger.addHandler(logging.NullHandler())
    return logger


# ============================================================================
# Normalization
# ============================================================================

class RecordNormalizer:
    """Maps real Excel headers to canonical concepts and cleans cell values.

    The original header text is always kept for display.
    """

    ITEM_NUMBER_ALIASES = [
        "item number", "item", "item no", "item nr", "item num", "part number", "part no",
        "part nr", "part num", "pn", "p n", "no de parte", "numero de parte",
        "num de parte", "component", "componente", "child", "child item",
    ]
    ALTERNATIVE_ALIASES = ["alternative", "alternativa", "alternate", "alt", "alt no", "alternative no"]
    MANUFACTURER_ALIASES = [
        "manufacturer", "manufacturer name", "mfr", "mfg", "mfr name", "vendor", "fabricante",
    ]
    # "type": in this workbook the AVL table header "Type" (column D) is pasted by
    # the macro Fetch_Details into the Start sheet column titled "MPN".
    MPN_ALIASES = [
        "mpn", "manufacturer part number", "manufacturer pn", "manufacturer p n", "mfr pn",
        "mfr part number", "mfg pn", "mfg part number", "type",
    ]
    TEXT_ALIASES = ["text", "description", "desc", "descripcion", "item description", "item text"]
    OBSOLETE_ALIASES = ["obsolete", "obsoleto"]
    # "stopped": the AVL header "Stopped" (column G) is pasted into the Start
    # column titled "Status"; the macro treats "Yes" as "Don't use".
    STATUS_ALIASES = ["status", "stopped", "estatus", "estado"]
    LEAD_FREE_ALIASES = ["lead free", "leadfree", "pb free", "rohs"]
    RECEIVED_LEAD_FREE_ALIASES = ["received lead free", "received leadfree", "received pb free"]
    VALID_UNTIL_ALIASES = ["valid until", "valid to", "expiration", "expiration date", "vigencia"]
    TEMPERATURE_ALIASES = ["temperature", "temp", "temperatura", "peak temperature", "max temperature"]
    MARKING_ALIASES = ["marking", "marcado", "marking code", "top marking", "part marking"]
    DURATION_ALIASES = ["duration", "duracion"]
    # Hierarchy concepts (only used when populated in the workbook)
    PARENT_ALIASES = [
        "parent", "parent item", "parent part", "parent pn", "parent part number",
        "assembly", "top level", "top level assembly", "padre", "ensamble",
    ]
    LEVEL_ALIASES = ["level", "bom level", "nivel", "lvl"]
    QUANTITY_ALIASES = ["quantity", "qty", "cantidad", "qty per"]

    FIELDS = OrderedDict([
        ("item_number", ("ITEM NUMBER", ITEM_NUMBER_ALIASES)),
        ("alternative", ("ALTERNATIVE", ALTERNATIVE_ALIASES)),
        ("manufacturer", ("MANUFACTURER", MANUFACTURER_ALIASES)),
        ("mpn", ("MPN", MPN_ALIASES)),
        ("text", ("DESCRIPTION", TEXT_ALIASES)),
        ("obsolete", ("OBSOLETE", OBSOLETE_ALIASES)),
        ("status", ("STATUS", STATUS_ALIASES)),
        ("lead_free", ("LEAD FREE", LEAD_FREE_ALIASES)),
        ("received_lead_free", ("RECEIVED LEAD FREE", RECEIVED_LEAD_FREE_ALIASES)),
        ("valid_until", ("VALID UNTIL", VALID_UNTIL_ALIASES)),
        ("temperature", ("TEMPERATURE", TEMPERATURE_ALIASES)),
        ("marking", ("MARKING", MARKING_ALIASES)),
        ("duration", ("DURATION", DURATION_ALIASES)),
        ("parent", ("ASSEMBLY / PARENT", PARENT_ALIASES)),
        ("level", ("BOM LEVEL", LEVEL_ALIASES)),
        ("quantity", ("QUANTITY", QUANTITY_ALIASES)),
    ])

    NULL_TOKENS = {"none", "nan", "null", "#n/a", "n/a", "#ref!", "#value!", "#name?", "nat"}
    YES_TOKENS = {"yes", "y", "si", "sí", "true", "x"}
    NO_TOKENS = {"no", "n", "false"}

    def __init__(self, extra_aliases: dict | None = None):
        self._alias_map: dict[str, str] = {}
        for key, (_label, aliases) in self.FIELDS.items():
            for alias in aliases:
                self._alias_map.setdefault(self.normalize_header(alias), key)
        for key, aliases in (extra_aliases or {}).items():
            if key in self.FIELDS and isinstance(aliases, list):
                for alias in aliases:
                    self._alias_map[self.normalize_header(str(alias))] = key

    # -- headers -------------------------------------------------------------
    @staticmethod
    def normalize_header(text) -> str:
        s = str(text or "").strip().lower()
        s = (s.replace("á", "a").replace("é", "e").replace("í", "i")
             .replace("ó", "o").replace("ú", "u").replace("ñ", "n"))
        s = re.sub(r"[_\-./\\:#()\[\]]+", " ", s)
        s = re.sub(r"[^a-z0-9 ]+", "", s)
        return re.sub(r"\s+", " ", s).strip()

    def canonical_for(self, header) -> str | None:
        return self._alias_map.get(self.normalize_header(header))

    @classmethod
    def label_for(cls, key: str) -> str:
        return cls.FIELDS[key][0] if key in cls.FIELDS else key.upper()

    # -- values --------------------------------------------------------------
    @classmethod
    def clean_value(cls, value):
        """Return a display string, or None when the cell is empty."""
        if value is None:
            return None
        if isinstance(value, bool):
            return "Yes" if value else "No"
        if isinstance(value, float):
            if math.isnan(value) or math.isinf(value):
                return None
            if value.is_integer():
                return str(int(value))
            return repr(value) if len(repr(value)) <= 15 else f"{value:.10g}"
        if isinstance(value, int):
            return str(value)
        if isinstance(value, dt.datetime):
            if value.time() == dt.time(0, 0):
                return value.strftime("%Y-%m-%d")
            return value.strftime("%Y-%m-%d %H:%M")
        if isinstance(value, dt.date):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, dt.time):
            return value.strftime("%H:%M:%S")
        text = str(value).replace("\r\n", "\n").replace("\r", "\n")
        # Strip the outer whitespace only; inner spaces and line breaks are kept
        # (markings can contain meaningful double spaces / line breaks).
        text = "\n".join(line.rstrip() for line in text.split("\n")).strip()
        if not text or text.lower() in cls.NULL_TOKENS:
            return None
        return text

    @staticmethod
    def display(value) -> str:
        return EMPTY if value is None or value == "" else str(value)

    @staticmethod
    def search_key(value) -> str:
        """Case-insensitive key with collapsed whitespace (scanner friendly)."""
        s = "".join(ch for ch in str(value or "") if ch.isprintable())
        return re.sub(r"\s+", " ", s).strip().casefold()

    @staticmethod
    def compact_key(value) -> str:
        return re.sub(r"\s+", "", str(value or "")).casefold()

    @classmethod
    def flag(cls, value) -> str | None:
        """'YES' / 'NO' / 'OTHER' / None for yes-no style columns."""
        if value is None:
            return None
        v = str(value).strip().casefold()
        if v in cls.YES_TOKENS:
            return "YES"
        if v in cls.NO_TOKENS:
            return "NO"
        return "OTHER"


# ============================================================================
# Data model
# ============================================================================

class Record:
    """One real data row of the workbook. `values` keeps the ORIGINAL row."""

    __slots__ = ("sheet", "row", "values", "fields", "headers_by_key", "uid", "has_detail", "order")

    def __init__(self, sheet: str, row: int, values: "OrderedDict[str, str | None]",
                 fields: dict, headers_by_key: dict, order: int):
        self.sheet = sheet
        self.row = row
        self.values = values                  # original header -> cleaned value
        self.fields = fields                  # canonical key -> cleaned value
        self.headers_by_key = headers_by_key  # canonical key -> original header
        self.uid = f"{sheet}!{row}"
        self.order = order
        self.has_detail = any(v is not None for k, v in fields.items() if k != "item_number") or any(
            v is not None for h, v in values.items()
            if headers_by_key.get("item_number") != h and h not in headers_by_key.values())

    def get(self, key: str):
        return self.fields.get(key)

    def status_flag(self) -> str | None:
        """Macro rule (Fetch_Details / FETCH_EMP_DETAILS): Status/Stopped = "Yes" -> Don't use."""
        return RecordNormalizer.flag(self.fields.get("status"))

    def __repr__(self):  # pragma: no cover - debug helper
        return f"<Record {self.uid} {self.fields.get('item_number')}>"


class SheetTable:
    def __init__(self, name: str):
        self.name = name
        self.header_row: int | None = None
        self.columns: list[tuple[int, str]] = []    # (column index, original header)
        self.key_by_header: dict[str, str] = {}     # original header -> canonical key
        self.header_by_key: dict[str, str] = {}     # canonical key -> original header
        self.ignored_columns: list[str] = []        # header cells outside the data block
        self.records: list[Record] = []
        self.preamble: list[tuple[int, list[str]]] = []  # non-empty rows above the header
        self.scanned_rows = 0
        self.structured = False

    @property
    def headers(self) -> list[str]:
        return [h for _, h in self.columns]


# ============================================================================
# Workbook readers
# ============================================================================

def col_letters_to_index(letters: str) -> int:
    idx = 0
    for ch in letters:
        idx = idx * 26 + (ord(ch.upper()) - 64)
    return idx - 1


class BuiltinXlsxReader:
    """Minimal read-only XLSX/XLSM reader: zipfile + xml.etree only."""

    name = "builtin (zipfile + xml)"
    _DATE_BUILTIN_FMTS = set(range(14, 23)) | {27, 30, 36, 45, 46, 47, 50, 57}

    def __init__(self, path: Path):
        self.path = Path(path)
        self._zip = zipfile.ZipFile(self.path, "r")
        self._names = set(self._zip.namelist())
        self._shared: list[str] = []
        self._date_styles: set[int] = set()
        self._date1904 = False
        self._sheets: "OrderedDict[str, str]" = OrderedDict()
        self._read_workbook()
        self._read_shared_strings()
        self._read_styles()

    @staticmethod
    def _local(tag: str) -> str:
        return tag.rsplit("}", 1)[-1]

    def _xml(self, member: str):
        with self._zip.open(member) as fh:
            return ET.parse(fh).getroot()

    def _read_workbook(self):
        root = self._xml("xl/workbook.xml")
        rels = {}
        rel_member = "xl/_rels/workbook.xml.rels"
        if rel_member in self._names:
            for rel in self._xml(rel_member):
                target = rel.get("Target", "")
                if target.startswith("/"):
                    target = target.lstrip("/")
                elif not target.startswith("xl/"):
                    target = "xl/" + target
                rels[rel.get("Id")] = os.path.normpath(target).replace("\\", "/")
        for el in root.iter():
            tag = self._local(el.tag)
            if tag == "workbookPr" and el.get("date1904") in ("1", "true"):
                self._date1904 = True
            if tag == "sheet":
                rid = next((v for k, v in el.attrib.items() if self._local(k) == "id"), None)
                target = rels.get(rid)
                if target and target in self._names:
                    self._sheets[el.get("name")] = target

    def _read_shared_strings(self):
        member = "xl/sharedStrings.xml"
        if member not in self._names:
            return
        with self._zip.open(member) as fh:
            for _event, el in ET.iterparse(fh, events=("end",)):
                if self._local(el.tag) == "si":
                    parts = []
                    for sub in el.iter():
                        t = self._local(sub.tag)
                        if t == "t":
                            parts.append(sub.text or "")
                        elif t == "rPh":  # phonetic runs are not cell text
                            break
                    self._shared.append("".join(parts))
                    el.clear()

    def _read_styles(self):
        member = "xl/styles.xml"
        if member not in self._names:
            return
        root = self._xml(member)
        custom_dates = set()
        for el in root.iter():
            if self._local(el.tag) == "numFmt":
                code = re.sub(r'"[^"]*"|\[[^\]]*\]', "", el.get("formatCode", "").lower())
                if re.search(r"[dmyh]", code) and "general" not in code:
                    custom_dates.add(int(el.get("numFmtId", "0")))
        for el in root.iter():
            if self._local(el.tag) == "cellXfs":
                for i, xf in enumerate([x for x in el if self._local(x.tag) == "xf"]):
                    fid = int(xf.get("numFmtId", "0"))
                    if fid in self._DATE_BUILTIN_FMTS or fid in custom_dates:
                        self._date_styles.add(i)
                break

    def sheet_names(self) -> list[str]:
        return list(self._sheets.keys())

    def _convert_serial(self, serial: float):
        base = dt.datetime(1904, 1, 1) if self._date1904 else dt.datetime(1899, 12, 30)
        try:
            return base + dt.timedelta(days=serial)
        except (OverflowError, ValueError):
            return serial

    def iter_rows(self, sheet_name: str):
        """Yield (excel_row_number, [values...])."""
        member = self._sheets[sheet_name]
        row_counter = 0
        with self._zip.open(member) as fh:
            for _event, el in ET.iterparse(fh, events=("end",)):
                if self._local(el.tag) != "row":
                    continue
                r_attr = el.get("r")
                row_counter = int(r_attr) if r_attr and r_attr.isdigit() else row_counter + 1
                cells = {}
                next_col = 0
                for c in el:
                    if self._local(c.tag) != "c":
                        continue
                    ref = c.get("r")
                    if ref:
                        m = re.match(r"([A-Za-z]+)", ref)
                        col = col_letters_to_index(m.group(1)) if m else next_col
                    else:
                        col = next_col
                    next_col = col + 1
                    ctype = c.get("t", "n")
                    raw = None
                    inline = []
                    for sub in c:
                        t = self._local(sub.tag)
                        if t == "v":
                            raw = sub.text
                        elif t == "is":
                            inline.extend(x.text or "" for x in sub.iter() if self._local(x.tag) == "t")
                    if ctype == "inlineStr":
                        value = "".join(inline)
                    elif raw is None:
                        continue
                    elif ctype == "s":
                        try:
                            value = self._shared[int(raw)]
                        except (ValueError, IndexError):
                            value = None
                    elif ctype == "b":
                        value = raw.strip() == "1"
                    elif ctype in ("str", "e"):
                        value = raw
                    else:
                        try:
                            num = float(raw)
                        except ValueError:
                            value = raw
                        else:
                            style = int(c.get("s", "0") or 0)
                            value = self._convert_serial(num) if style in self._date_styles else num
                    if value is not None and value != "":
                        cells[col] = value
                el.clear()
                if cells:
                    width = max(cells) + 1
                    yield row_counter, [cells.get(i) for i in range(width)]
                else:
                    yield row_counter, []

    def close(self):
        self._zip.close()


class OpenpyxlReader:
    name = "openpyxl"

    def __init__(self, path: Path):
        # read_only + data_only: cached values, nothing is ever written back.
        self._wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True, keep_vba=True)

    def sheet_names(self) -> list[str]:
        return list(self._wb.sheetnames)

    def iter_rows(self, sheet_name: str):
        ws = self._wb[sheet_name]
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()  # do not trust a stale <dimension> tag
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            yield idx, list(row)

    def close(self):
        try:
            self._wb.close()
        except Exception as exc:
            logging.getLogger("gpv").debug("openpyxl close: %s", exc)


# ============================================================================
# Excel data provider
# ============================================================================

class ExcelLoadError(Exception):
    """Error with a short operator-friendly message."""


class ExcelDataProvider:
    HEADER_SCAN_ROWS = 60

    def __init__(self, excel_path: Path, normalizer: RecordNormalizer, reader: str = "auto",
                 logger: logging.Logger | None = None):
        self.excel_path = Path(excel_path)
        self.normalizer = normalizer
        self.reader_preference = (reader or "auto").lower()
        self.log = logger or logging.getLogger("gpv")
        self.tables: "OrderedDict[str, SheetTable]" = OrderedDict()
        self.records: list[Record] = []
        self.reader_used = ""
        self.loaded_at: dt.datetime | None = None
        self.file_mtime: dt.datetime | None = None
        self.load_seconds = 0.0
        self.all_sheet_names: list[str] = []
        self._item_index: dict[str, list[Record]] = {}
        self._item_compact: dict[str, list[Record]] = {}
        self._field_index: dict[str, dict[str, list[Record]]] = {}
        self._search_rows: list[tuple[str, str, str, Record]] = []
        self.hierarchy_mode = "logical"
        self.children_map: dict[str, list[tuple[Record, str]]] = {}

    # -- loading ---------------------------------------------------------------
    def load_workbook(self) -> "ExcelDataProvider":
        t0 = time.perf_counter()
        path = self.excel_path
        if not path.exists():
            raise ExcelLoadError(f"Excel file not found:\n{self._rel(path)}")
        if path.suffix.lower() not in (".xlsx", ".xlsm"):
            raise ExcelLoadError(f"Unsupported file type: {path.suffix} (use .xlsx or .xlsm)")
        tmpdir = tempfile.mkdtemp(prefix="gpv_pc_")
        snapshot = Path(tmpdir) / ("snapshot" + path.suffix.lower())
        try:
            try:
                shutil.copy2(path, snapshot)  # never read/lock the original
            except PermissionError:
                raise ExcelLoadError("The Excel file is locked by another program.\nClose it or try RELOAD (F5).")
            except OSError as exc:
                raise ExcelLoadError(f"Excel file could not be copied for reading ({exc.__class__.__name__}).")
            self.file_mtime = dt.datetime.fromtimestamp(path.stat().st_mtime)
            reader = self._open_reader(snapshot)
            try:
                self._parse(reader)
            finally:
                reader.close()
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)
        self._build_indexes()
        self.loaded_at = dt.datetime.now()
        self.load_seconds = time.perf_counter() - t0
        for t in self.tables.values():
            self.log.info("Sheet '%s': header row %s, %d columns, %d records%s", t.name, t.header_row,
                          len(t.columns), len(t.records), "" if t.structured else " (no recognised header)")
        self.log.info("Excel loaded: %s | reader=%s | sheets=%s | rows=%d | %.2fs",
                      path.name, self.reader_used, self.all_sheet_names, len(self.records), self.load_seconds)
        return self

    def reload(self) -> "ExcelDataProvider":
        self.log.info("Reload requested")
        return self.load_workbook()

    def _rel(self, path: Path) -> str:
        try:
            return str(path.resolve().relative_to(BASE_DIR))
        except ValueError:
            return str(path)

    def _open_reader(self, snapshot: Path):
        pref = self.reader_preference
        if pref in ("auto", "openpyxl") and openpyxl is not None:
            try:
                reader = OpenpyxlReader(snapshot)
                self.reader_used = reader.name
                return reader
            except Exception as exc:
                self.log.warning("openpyxl could not open the workbook (%s); using builtin reader", exc)
        try:
            reader = BuiltinXlsxReader(snapshot)
        except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
            self.log.error("Workbook is not a valid xlsx/xlsm: %s", exc)
            raise ExcelLoadError("The Excel file is damaged or is being saved.\nTry RELOAD (F5) in a moment.")
        self.reader_used = reader.name
        return reader

    def _parse(self, reader):
        tables: "OrderedDict[str, SheetTable]" = OrderedDict()
        records: list[Record] = []
        self.all_sheet_names = reader.sheet_names()
        order = 0
        for name in self.all_sheet_names:
            table = SheetTable(name)
            rows = reader.iter_rows(name)
            buffer: list[tuple[int, list]] = []
            for rownum, values in rows:
                table.scanned_rows += 1
                if any(self.normalizer.clean_value(v) is not None for v in values):
                    buffer.append((rownum, values))
                    if len(buffer) >= self.HEADER_SCAN_ROWS:
                        break
            header_pos = self.detect_headers(buffer, table)
            if header_pos is None:
                # Consume the iterator to release the sheet; nothing structured here.
                for _ in rows:
                    table.scanned_rows += 1
                tables[name] = table
                continue

            def all_rows():
                yield from buffer[header_pos + 1:]
                for item in rows:
                    table.scanned_rows += 1
                    yield item

            for rownum, values in all_rows():
                rec = self._make_record(table, rownum, values, order)
                if rec is not None:
                    table.records.append(rec)
                    records.append(rec)
                    order += 1
            tables[name] = table
        self.tables = tables
        self.records = records

    def detect_headers(self, buffer, table: SheetTable):
        """Find the header row by alias score; keep the contiguous header block
        that contains recognised columns (ignores side legends such as
        'Colores / Significado')."""
        best_pos, best_score = None, 0
        for pos, (_rownum, values) in enumerate(buffer):
            score = sum(1 for v in values
                        if isinstance(v, str) and self.normalizer.canonical_for(v))
            if score > best_score:
                best_pos, best_score = pos, score
        if best_pos is None or best_score < 2:
            return None
        rownum, values = buffer[best_pos]
        texts = [self.normalizer.clean_value(v) for v in values]
        blocks, current = [], []
        for idx, t in enumerate(texts):
            if t is None:
                if current:
                    blocks.append(current)
                current = []
            else:
                current.append((idx, t))
        if current:
            blocks.append(current)
        used = []
        for block in blocks:
            if any(self.normalizer.canonical_for(h) for _, h in block):
                used.extend(block)
            else:
                table.ignored_columns.extend(h for _, h in block)
        seen = set()
        for idx, header in used:
            label = header if header not in seen else f"{header} ({idx + 1})"
            seen.add(label)
            table.columns.append((idx, label))
            key = self.normalizer.canonical_for(header)
            if key and key not in table.header_by_key:
                table.header_by_key[key] = label
                table.key_by_header[label] = key
        table.header_row = rownum
        table.structured = True
        table.preamble = [(r, [self.normalizer.display(self.normalizer.clean_value(v)) for v in vals
                               if self.normalizer.clean_value(v) is not None])
                          for r, vals in buffer[:best_pos]]
        return best_pos

    def _make_record(self, table: SheetTable, rownum: int, values: list, order: int):
        data = OrderedDict()
        filled = 0
        for idx, header in table.columns:
            v = self.normalizer.clean_value(values[idx] if idx < len(values) else None)
            data[header] = v
            if v is not None:
                filled += 1
        if not filled:
            return None  # completely empty row
        if all(data[h] is not None and self.normalizer.normalize_header(data[h]) ==
               self.normalizer.normalize_header(h) for h in data):
            return None  # repeated header row
        fields = {key: data[h] for h, key in table.key_by_header.items()}
        return Record(table.name, rownum, data, fields, table.header_by_key, order)

    def _build_indexes(self):
        item_idx, item_compact, field_idx, search_rows = {}, {}, {}, []
        nz = self.normalizer
        searchable = ("item_number", "parent", "mpn")
        for rec in self.records:
            item = rec.get("item_number")
            if item:
                item_idx.setdefault(nz.search_key(item), []).append(rec)
                item_compact.setdefault(nz.compact_key(item), []).append(rec)
            for key in searchable:
                v = rec.get(key)
                if v:
                    k = nz.search_key(v)
                    field_idx.setdefault(key, {}).setdefault(k, []).append(rec)
                    search_rows.append((key, k, nz.compact_key(v), rec))
        self._item_index, self._item_compact = item_idx, item_compact
        self._field_index, self._search_rows = field_idx, search_rows
        self._build_hierarchy()

    def _build_hierarchy(self):
        """Parent/child only from real columns: 'parent' values or numeric 'level'."""
        nz = self.normalizer
        children: dict[str, list[tuple[Record, str]]] = {}
        mode = "logical"
        if any(r.get("parent") and r.get("item_number") for r in self.records):
            for r in self.records:
                p, c = r.get("parent"), r.get("item_number")
                if p and c and nz.search_key(p) != nz.search_key(c):
                    children.setdefault(nz.search_key(p), []).append((r, "Parent column"))
            mode = "excel-parent"
        elif any(self._level_of(r) is not None for r in self.records):
            for table in self.tables.values():
                stack: list[tuple[int, str]] = []
                for r in table.records:
                    lvl, item = self._level_of(r), r.get("item_number")
                    if lvl is None or not item:
                        continue
                    while stack and stack[-1][0] >= lvl:
                        stack.pop()
                    if stack:
                        children.setdefault(stack[-1][1], []).append((r, "Level column"))
                    stack.append((lvl, nz.search_key(item)))
            mode = "excel-level"
        self.children_map = children
        self.hierarchy_mode = mode if children else "logical"

    @staticmethod
    def _level_of(rec: Record):
        v = rec.get("level")
        if v is None:
            return None
        m = re.search(r"-?\d+", v.replace(".", ""))
        return int(m.group()) if m and v.strip(". ").replace(".", "").lstrip("-").isdigit() else None

    # -- queries -----------------------------------------------------------------
    def list_sheets(self) -> list[str]:
        return list(self.all_sheet_names)

    def get_records(self, sheet: str | None = None) -> list[Record]:
        if sheet and sheet in self.tables:
            return list(self.tables[sheet].records)
        return list(self.records)

    def find_item(self, value) -> list[Record]:
        nz = self.normalizer
        hits = self._item_index.get(nz.search_key(value))
        if not hits:
            hits = self._item_compact.get(nz.compact_key(value), [])
        return list(hits)

    def find_mpn(self, value) -> list[Record]:
        return list(self._field_index.get("mpn", {}).get(self.normalizer.search_key(value), []))

    def find_assembly(self, value) -> list[Record]:
        """Rows whose Assembly/Parent column matches; falls back to Item number
        because this workbook has no assembly column."""
        hits = self._field_index.get("parent", {}).get(self.normalizer.search_key(value), [])
        return list(hits) if hits else self.find_item(value)

    def get_alternatives(self, item_number) -> list[Record]:
        recs = self.find_item(item_number)
        return sorted(recs, key=lambda r: (not r.has_detail, self._sheet_pos(r.sheet),
                                           natural_key(r.get("alternative") or ""), r.row))

    def get_record_details(self, record: Record) -> list[tuple[str, str]]:
        rows = [("Sheet", record.sheet), ("Excel row", str(record.row))]
        rows += [(h, self.normalizer.display(v)) for h, v in record.values.items()]
        return rows

    def _sheet_pos(self, name: str) -> int:
        try:
            return self.all_sheet_names.index(name)
        except ValueError:
            return 999

    # -- metadata ----------------------------------------------------------------
    def has_field(self, key: str) -> bool:
        return any(key in t.header_by_key for t in self.tables.values())

    def populated_count(self, key: str) -> int:
        return sum(1 for r in self.records if r.get(key) is not None)

    def distinct_values(self, key: str, limit: int = 500) -> list[str]:
        vals = {r.get(key) for r in self.records if r.get(key) is not None}
        return sorted(vals, key=natural_key)[:limit]

    def sheet_summary(self) -> str:
        parts = [f"{t.name} ({len(t.records):,})" for t in self.tables.values() if t.structured]
        return " | ".join(parts) if parts else EMPTY

    def item_count(self) -> int:
        return len(self._item_index)


def natural_key(text):
    return [int(t) if t.isdigit() else t.casefold() for t in re.split(r"(\d+)", str(text))]


# ============================================================================
# Search
# ============================================================================

class SearchResult:
    def __init__(self, query: str):
        self.query = query
        self.items: "OrderedDict[str, dict]" = OrderedDict()  # item key -> info
        self.match_kind = ""        # exact | prefix | contains | none
        self.truncated = 0
        self.elapsed_ms = 0.0

    @property
    def found(self) -> bool:
        return bool(self.items)

    def matched_uids(self) -> set:
        """Rows to highlight: only those found through MPN / Assembly columns."""
        return {r.uid for info in self.items.values() if info["field"] != "item_number" for r in info["matched"]}


class SearchEngine:
    FIELD_LABELS = {"item_number": "Item number", "parent": "Assembly", "mpn": "MPN"}

    def __init__(self, provider: ExcelDataProvider, max_items: int = 150):
        self.provider = provider
        self.max_items = max_items

    def search(self, query: str) -> SearchResult:
        t0 = time.perf_counter()
        nz = self.provider.normalizer
        res = SearchResult(query)
        key, compact = nz.search_key(query), nz.compact_key(query)
        if not key:
            res.match_kind = "none"
            return res
        hits: list[tuple[int, str, Record]] = []
        # 1) exact (case/space insensitive) - has absolute priority
        for rank, (field, k, c, rec) in enumerate(self.provider._search_rows):
            if k == key or c == compact:
                hits.append((0, field, rec))
        kind = "exact"
        if not hits and len(compact) >= 2:
            for field, k, c, rec in self.provider._search_rows:
                if c.startswith(compact):
                    hits.append((1, field, rec))
            kind = "prefix"
        if not hits and len(compact) >= 3:
            for field, k, c, rec in self.provider._search_rows:
                if compact in c:
                    hits.append((2, field, rec))
            kind = "contains"
        res.match_kind = kind if hits else "none"
        field_order = {"item_number": 0, "parent": 1, "mpn": 2}
        hits.sort(key=lambda h: (h[0], field_order.get(h[1], 9), len(h[2].get(h[1]) or ""), h[2].order))
        for _rank, field, rec in hits:
            item = rec.get("item_number") or "(no item number)"
            if field == "parent" and rec.get("parent"):
                item = rec.get("parent")
            ikey = nz.search_key(item)
            if ikey not in res.items:
                if len(res.items) >= self.max_items:
                    res.truncated += 1
                    continue
                res.items[ikey] = {"item": item, "field": field, "matched": []}
            if rec not in res.items[ikey]["matched"]:
                res.items[ikey]["matched"].append(rec)
        res.elapsed_ms = (time.perf_counter() - t0) * 1000
        return res


# ============================================================================
# BOM tree building
# ============================================================================

class TreeNode:
    __slots__ = ("text", "kind", "record", "item_key", "info", "tags", "children", "open")

    def __init__(self, text, kind, record=None, item_key=None, info="", tags=(), open_=False):
        self.text, self.kind, self.record, self.item_key = text, kind, record, item_key
        self.info, self.tags, self.children, self.open = info, tuple(tags), [], open_

    def add(self, node: "TreeNode") -> "TreeNode":
        self.children.append(node)
        return node


class BomBuilder:
    """Builds the tree shown in BOM STRUCTURE.

    * excel-parent / excel-level: real hierarchy from populated workbook columns.
    * logical: Item -> Alternatives -> Manufacturer / MPN / Marking (LOGICAL VIEW,
      not a FactoryLogix genealogy).
    """

    MAX_DEPTH = 25

    def __init__(self, provider: ExcelDataProvider):
        self.provider = provider

    @property
    def view_name(self) -> str:
        return {"excel-parent": "EXCEL BOM (PARENT)", "excel-level": "EXCEL BOM (LEVEL)"}.get(
            self.provider.hierarchy_mode, "LOGICAL VIEW")

    def build(self, result: SearchResult, mode: str) -> TreeNode:
        matched = result.matched_uids()
        items = list(result.items.items())
        if len(items) == 1:
            ikey, info = items[0]
            root = self._item_node(ikey, info["item"], mode, matched, depth=0, path=set(), icon="\U0001F4E6")
            root.open = True
            return root
        root = TreeNode(f"\U0001F4E6  SEARCH: {result.query}   ({len(items)} items)", "search",
                        info=f"{result.match_kind} match", open_=True)
        for ikey, info in items:
            node = self._item_node(ikey, info["item"], mode, matched, depth=1, path=set(), icon="▣")
            node.open = len(items) <= 3
            root.add(node)
        if result.truncated:
            root.add(TreeNode(f"… {result.truncated} more matching rows not listed - refine the search",
                              "info", tags=("muted",)))
        return root

    def _item_node(self, ikey, item_display, mode, matched, depth, path, icon, via="") -> TreeNode:
        prov = self.provider
        recs = prov.get_alternatives(item_display)
        desc = next((r.get("text") for r in recs if r.get("text")), None)
        label = f"{icon}  {item_display}" + (f"  -  {desc}" if desc else "")
        details = [r for r in recs if r.has_detail]
        refs = [r for r in recs if not r.has_detail]
        info = f"{len(details)} alt." if details else (f"{len(refs)} row(s), no detail" if refs else "")
        if via:
            info = f"{via}  {info}".strip()
        tags = ["item"]
        if any(r.uid in matched for r in recs):
            tags.append("match")
        node = TreeNode(label, "item", item_key=ikey, info=info, tags=tags, open_=depth <= 1)

        hidden = 0
        for rec in details:
            flag = rec.status_flag()
            if mode == "PRODUCTION" and flag == "YES":
                hidden += 1
                continue
            alt_node = node.add(self._alternative_node(rec, mode, matched))
            alt_node.open = depth == 0 and len(details) <= 6
        if hidden:
            node.add(TreeNode(f"⊘  {hidden} alternative(s) hidden in PRODUCTION (Status = Yes)", "info",
                              tags=("muted",)))
        if refs:
            by_sheet: "OrderedDict[str, list[Record]]" = OrderedDict()
            for r in refs:
                by_sheet.setdefault(r.sheet, []).append(r)
            for sheet, rows in by_sheet.items():
                if mode == "PRODUCTION":
                    node.add(TreeNode(f"▫  {len(rows)} row(s) in '{sheet}' without detail data", "info",
                                      item_key=ikey, tags=("muted",)))
                    continue
                grp = node.add(TreeNode(f"▫  '{sheet}' rows without detail data ({len(rows)})", "refgroup",
                                        item_key=ikey, info="Item number only", tags=("muted",),
                                        open_=not details))
                for r in rows:
                    grp.add(TreeNode(f"Row {r.row}", "ref", record=r, info=f"{r.sheet}!{r.row}",
                                     tags=("muted", "match") if r.uid in matched else ("muted",)))

        # Real hierarchy children (only when the workbook provides it)
        if prov.hierarchy_mode != "logical" and depth < self.MAX_DEPTH and ikey not in path:
            seen_child = set()
            for child_rec, source in prov.children_map.get(ikey, []):
                child_item = child_rec.get("item_number")
                ck = prov.normalizer.search_key(child_item)
                if ck in seen_child:
                    continue
                seen_child.add(ck)
                qty = child_rec.get("quantity")
                via_txt = f"Qty {qty}" if qty else ""
                node.add(self._item_node(ck, child_item, mode, matched, depth + 1, path | {ikey},
                                         icon="▣", via=via_txt))
        return node

    def _alternative_node(self, rec: Record, mode: str, matched: set) -> TreeNode:
        nz = RecordNormalizer
        alt = rec.get("alternative")
        parts = [f"ALT {alt}" if alt else f"ROW {rec.row}"]
        if rec.get("manufacturer"):
            parts.append(rec.get("manufacturer"))
        if rec.get("mpn"):
            parts.append(rec.get("mpn"))
        flag = rec.status_flag()
        use_txt = {"YES": "DON'T USE", "NO": "USE"}.get(flag, "")
        tags = ["alt"]
        if flag == "YES":
            tags.append("dontuse")
        elif flag == "NO":
            tags.append("use")
        if nz.flag(rec.get("obsolete")) == "YES":
            tags.append("obsolete")
            use_txt = (use_txt + "  OBSOLETE").strip()
        if rec.uid in matched:
            tags.append("match")
        info = use_txt if mode == "PRODUCTION" else f"{use_txt}   {rec.sheet}!{rec.row}".strip()
        node = TreeNode("◆  " + "  ·  ".join(parts), "alt", record=rec, info=info, tags=tags)
        for key in ("manufacturer", "mpn", "marking"):
            if key not in rec.headers_by_key:
                continue
            value = rec.get(key)
            shown = nz.display(value).replace("\n", "  ⏎  ")
            node.add(TreeNode(f"{nz.label_for(key).title() if key != 'mpn' else 'MPN'}: {shown}", "attr",
                              record=rec, tags=("attr",)))
        return node


# ============================================================================
# UI helpers
# ============================================================================

def pick_font(candidates, fallback="TkDefaultFont"):
    try:
        families = set(tkfont.families())
    except Exception:
        return fallback
    for c in candidates:
        if c in families:
            return c
    return fallback


class InfoPanel(tk.Frame if tk else object):
    """White panel with soft border, title and label/value rows."""

    def __init__(self, master, title, fonts, columns=1, title_color=C_NAVY):
        super().__init__(master, bg=C_PANEL, highlightthickness=1, highlightbackground=C_BORDER,
                         highlightcolor=C_BORDER, bd=0)
        self.fonts = fonts
        self.columns = columns
        self.title_lbl = tk.Label(self, text=title, bg=C_PANEL, fg=title_color, font=fonts["panel"], anchor="w")
        self.title_lbl.pack(fill="x", padx=10, pady=(7, 3))
        self.body = tk.Frame(self, bg=C_PANEL)
        self.body.pack(fill="both", expand=True, padx=10, pady=(0, 8))
        self._values: list[tk.Label] = []
        self._wrap = 300
        self.bind("<Configure>", self._on_resize)

    def set_rows(self, rows):
        """rows: list of (label, value, style). style in None|good|bad|warn|muted|strong."""
        for w in self.body.winfo_children():
            w.destroy()
        self._values = []
        for c in range(self.columns * 2):
            self.body.columnconfigure(c, weight=1 if c % 2 else 0, uniform="" if c % 2 else f"lbl{c}")
        if not rows:
            tk.Label(self.body, text=EMPTY, bg=C_PANEL, fg=C_MUTED, font=self.fonts["value"]).grid(
                row=0, column=0, sticky="w")
            return
        colors = {"good": C_OK, "bad": C_RED, "warn": C_AMBER, "muted": C_MUTED, "strong": C_NAVY}
        for i, (label, value, style) in enumerate(rows):
            r, c = divmod(i, self.columns)
            tk.Label(self.body, text=f"{label}:", bg=C_PANEL, fg=C_TEXT, font=self.fonts["label"],
                     anchor="nw", justify="left").grid(row=r, column=c * 2, sticky="nw", padx=(0, 8), pady=1)
            text = RecordNormalizer.display(value)
            fg = colors.get(style, C_TEXT) if text != EMPTY else C_MUTED
            fnt = self.fonts["value_b"] if style in ("good", "bad", "strong") else self.fonts["value"]
            v = tk.Label(self.body, text=text, bg=C_PANEL, fg=fg, font=fnt, anchor="nw", justify="left",
                         wraplength=self._wrap)
            v.grid(row=r, column=c * 2 + 1, sticky="nw", padx=(0, 14), pady=1)
            self._values.append(v)

    def _on_resize(self, event):
        wrap = max(140, int(event.width / self.columns) - 190)
        if abs(wrap - self._wrap) > 8:
            self._wrap = wrap
            for v in self._values:
                v.configure(wraplength=wrap)


class MarkingPanel(tk.Frame if tk else object):
    """EXPECTED MARKING - large monospace text, keeps spaces and line breaks."""

    def __init__(self, master, fonts):
        super().__init__(master, bg=C_PANEL, highlightthickness=1, highlightbackground=C_BORDER, bd=0)
        self.fonts = fonts
        head = tk.Frame(self, bg=C_PANEL)
        head.pack(fill="x", padx=10, pady=(7, 2))
        tk.Label(head, text="MARKING INFORMATION", bg=C_PANEL, fg=C_NAVY, font=fonts["panel"]).pack(side="left")
        self.source_lbl = tk.Label(head, text="", bg=C_PANEL, fg=C_MUTED, font=fonts["small"])
        self.source_lbl.pack(side="right")
        tk.Label(self, text="EXPECTED MARKING", bg=C_PANEL, fg=C_GREEN, font=fonts["label"], anchor="w").pack(
            fill="x", padx=10)
        self.text = tk.Text(self, height=2, bg="#F7F9FA", fg=C_TEXT, relief="flat", wrap="word",
                            font=fonts["marking"], highlightthickness=1, highlightbackground=C_SOFT,
                            padx=10, pady=6, cursor="arrow")
        self.text.pack(fill="x", padx=10, pady=(2, 10))
        self.text.tag_configure("key", font=fonts["small_b"], foreground=C_MUTED)
        self.text.tag_configure("empty", foreground=C_MUTED, font=fonts["value"])
        self.text.tag_configure("bad", foreground=C_RED)
        self.text.configure(state="disabled")

    def set_markings(self, entries, source=""):
        """entries: list of (prefix, marking or None, style)."""
        self.text.configure(state="normal")
        self.text.delete("1.0", "end")
        lines = 0
        if not entries:
            self.text.insert("end", "No marking data for the selection", "empty")
            lines = 1
        for i, (prefix, marking, style) in enumerate(entries):
            if i:
                self.text.insert("end", "\n")
            if prefix:
                self.text.insert("end", prefix + "\n", "key")
                lines += 1
            if marking is None:
                self.text.insert("end", f"{EMPTY}  (no marking in Excel)", "empty")
                lines += 1
            else:
                self.text.insert("end", marking, ("bad",) if style == "bad" else ())
                lines += marking.count("\n") + 1
        self.text.configure(height=max(2, min(lines, 14)), state="disabled")
        self.source_lbl.configure(text=source)


# ============================================================================
# Details window
# ============================================================================

class DetailsWindow(tk.Toplevel if tk else object):
    FILTER_KEYS = ["item_number", "manufacturer", "mpn", "alternative", "status", "obsolete"]
    COMBO_KEYS = {"manufacturer", "alternative", "status", "obsolete"}

    def __init__(self, app: "GPVProductConfigurationApp", item_filter: str = ""):
        super().__init__(app)
        self.app = app
        self.provider = app.provider
        self.fonts = app.fonts
        self.title(f"{APP_TITLE} - DETAILS")
        self.configure(bg=C_BG)
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f"{min(1250, sw - 60)}x{min(680, sh - 100)}")
        self.minsize(640, 380)
        self._sort_col, self._sort_rev = None, False
        self._after = None
        self._filtered: list[Record] = []
        self._build(item_filter)
        self.bind("<Escape>", lambda e: self.destroy())
        self.bind("<Control-c>", self._copy_selection)
        self.apply_filters()

    def _build(self, item_filter):
        f = self.fonts
        top = tk.Frame(self, bg=C_PANEL, highlightthickness=1, highlightbackground=C_BORDER)
        top.pack(fill="x", padx=10, pady=(10, 6))
        tk.Label(top, text="ALL COLUMNS FROM EXCEL", bg=C_PANEL, fg=C_NAVY, font=f["panel"]).grid(
            row=0, column=0, columnspan=8, sticky="w", padx=8, pady=(6, 4))
        tk.Label(top, text="Sheet", bg=C_PANEL, fg=C_TEXT, font=f["label"]).grid(row=1, column=0, sticky="w", padx=(8, 4))
        sheets = ["(All sheets)"] + [t.name for t in self.provider.tables.values() if t.structured]
        self.sheet_var = tk.StringVar(value=sheets[0])
        cb = ttk.Combobox(top, textvariable=self.sheet_var, values=sheets, state="readonly", width=16)
        cb.grid(row=1, column=1, sticky="w", padx=(0, 10), pady=2)
        cb.bind("<<ComboboxSelected>>", lambda e: self._rebuild_columns())
        tk.Label(top, text="Search", bg=C_PANEL, fg=C_TEXT, font=f["label"]).grid(row=1, column=2, sticky="w", padx=(0, 4))
        self.search_var = tk.StringVar()
        ttk.Entry(top, textvariable=self.search_var, width=28).grid(row=1, column=3, sticky="w", pady=2)
        self.search_var.trace_add("write", lambda *a: self._schedule())

        self.filter_vars: dict[str, tk.StringVar] = {}
        col = 0
        frow = 2
        for key in self.FILTER_KEYS:
            if not self.provider.has_field(key):
                continue  # only filters for columns that really exist
            var = tk.StringVar(value=item_filter if key == "item_number" else "")
            self.filter_vars[key] = var
            label = "MPN" if key == "mpn" else RecordNormalizer.label_for(key).title()
            tk.Label(top, text=label, bg=C_PANEL, fg=C_TEXT,
                     font=f["label"]).grid(row=frow, column=col, sticky="w", padx=(8 if col == 0 else 0, 4), pady=2)
            if key in self.COMBO_KEYS:
                values = [""] + self.provider.distinct_values(key)
                w = ttk.Combobox(top, textvariable=var, values=values, width=16)
                w.bind("<<ComboboxSelected>>", lambda e: self._schedule())
            else:
                w = ttk.Entry(top, textvariable=var, width=18)
            w.grid(row=frow, column=col + 1, sticky="w", padx=(0, 10), pady=2)
            var.trace_add("write", lambda *a: self._schedule())
            col += 2
            if col >= 8:
                col, frow = 0, frow + 1
        btns = tk.Frame(top, bg=C_PANEL)
        btns.grid(row=frow + 1, column=0, columnspan=8, sticky="we", padx=8, pady=(4, 6))
        ttk.Button(btns, text="CLEAR FILTERS", style="Tool.TButton", command=self._clear).pack(side="left")
        ttk.Button(btns, text="COPY ROWS (Ctrl+C)", style="Tool.TButton",
                   command=self._copy_selection).pack(side="left", padx=6)
        self.count_lbl = tk.Label(btns, text="", bg=C_PANEL, fg=C_MUTED, font=f["small"])
        self.count_lbl.pack(side="right")

        body = tk.Frame(self, bg=C_PANEL, highlightthickness=1, highlightbackground=C_BORDER)
        body.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        body.rowconfigure(0, weight=1)
        body.columnconfigure(0, weight=1)
        self.tree = ttk.Treeview(body, show="headings", selectmode="extended", style="Details.Treeview")
        vs = ttk.Scrollbar(body, orient="vertical", command=self.tree.yview)
        hs = ttk.Scrollbar(body, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vs.grid(row=0, column=1, sticky="ns")
        hs.grid(row=1, column=0, sticky="ew")
        self.tree.tag_configure("odd", background="#F6F8F9")
        self._rebuild_columns(apply=False)

    def _columns_for_sheet(self):
        """Returns list of (column id, heading, getter)."""
        sheet = self.sheet_var.get()
        cols = [("__sheet", "Sheet", lambda r: r.sheet), ("__row", "Row", lambda r: str(r.row))]
        tables = [t for t in self.provider.tables.values() if t.structured and
                  (sheet == "(All sheets)" or t.name == sheet)]
        merged: "OrderedDict[str, list[str]]" = OrderedDict()
        for t in tables:
            for h in t.headers:
                key = t.key_by_header.get(h) or "raw:" + RecordNormalizer.normalize_header(h)
                names = merged.setdefault(key, [])
                if h not in names:
                    names.append(h)
        for key, names in merged.items():
            heading = " / ".join(names)
            if key.startswith("raw:"):
                getter = (lambda k: lambda r: next((v for h, v in r.values.items()
                                                    if "raw:" + RecordNormalizer.normalize_header(h) == k), None))(key)
            else:
                getter = (lambda k: lambda r: r.get(k))(key)
            cols.append((key, heading, getter))
        return cols

    def _rebuild_columns(self, apply=True):
        self.columns = self._columns_for_sheet()
        ids = [c[0] for c in self.columns]
        self.tree.configure(columns=ids)
        for cid, heading, _g in self.columns:
            self.tree.heading(cid, text=heading, command=lambda c=cid: self._sort_by(c))
            width = 70 if cid in ("__sheet", "__row") else max(90, min(260, len(heading) * 9 + 30))
            if cid in ("text", "mpn", "marking"):
                width = 220
            self.tree.column(cid, width=width, minwidth=50, stretch=False, anchor="w")
        self._sort_col = None
        if apply:
            self.apply_filters()

    def _schedule(self):
        if self._after:
            self.after_cancel(self._after)
        self._after = self.after(250, self.apply_filters)

    def _clear(self):
        self.search_var.set("")
        for v in self.filter_vars.values():
            v.set("")

    def apply_filters(self):
        self._after = None
        nz = RecordNormalizer
        sheet = self.sheet_var.get()
        recs = self.provider.get_records(None if sheet == "(All sheets)" else sheet)
        conds = [(k, nz.search_key(v.get())) for k, v in self.filter_vars.items() if nz.search_key(v.get())]
        term = nz.search_key(self.search_var.get())
        out = []
        for r in recs:
            ok = True
            for k, needle in conds:
                val = r.get(k)
                if val is None or needle not in nz.search_key(val):
                    ok = False
                    break
            if ok and term:
                ok = any(v is not None and term in nz.search_key(v) for v in r.values.values())
            if ok:
                out.append(r)
        self._filtered = out
        if self._sort_col:
            self._sort_records()
        self._render()

    def _sort_by(self, cid):
        if self._sort_col == cid:
            self._sort_rev = not self._sort_rev
        else:
            self._sort_col, self._sort_rev = cid, False
        self._sort_records()
        self._render()

    def _sort_records(self):
        getter = next((g for c, _h, g in self.columns if c == self._sort_col), None)
        if getter is None:
            return

        def key(r):
            v = getter(r)
            if v is None:
                return (1, [])
            return (0, natural_key(v))
        present = [r for r in self._filtered if getter(r) is not None]
        missing = [r for r in self._filtered if getter(r) is None]
        present.sort(key=key, reverse=self._sort_rev)
        self._filtered = present + missing  # empty cells always at the end

    def _render(self):
        self.tree.delete(*self.tree.get_children())
        limit = int(self.app.settings.get("details_max_rows", 5000))
        getters = [g for _c, _h, g in self.columns]
        for i, r in enumerate(self._filtered[:limit]):
            vals = [RecordNormalizer.display(g(r)).replace("\n", " ⏎ ") for g in getters]
            self.tree.insert("", "end", iid=r.uid, values=vals, tags=("odd",) if i % 2 else ())
        for cid, heading, _g in self.columns:
            arrow = ""
            if cid == self._sort_col:
                arrow = "  ▼" if self._sort_rev else "  ▲"
            self.tree.heading(cid, text=heading + arrow)
        total = len(self._filtered)
        shown = min(total, limit)
        extra = f"  (showing first {shown:,} - refine filters)" if total > limit else ""
        self.count_lbl.configure(text=f"{total:,} matching rows{extra}")

    def _copy_selection(self, _event=None):
        sel = self.tree.selection()
        if not sel:
            return
        lines = ["\t".join(h for _c, h, _g in self.columns)]
        for iid in sel:
            lines.append("\t".join(str(v) for v in self.tree.item(iid, "values")))
        self.clipboard_clear()
        self.clipboard_append("\n".join(lines))
        self.count_lbl.configure(text=f"{len(sel)} row(s) copied to clipboard")


# ============================================================================
# Main application
# ============================================================================

class GPVProductConfigurationApp(tk.Tk if tk else object):
    MODES = ("ENGINEERING", "PRODUCTION")

    def __init__(self, settings: dict, logger: logging.Logger):
        super().__init__()
        self.settings = settings
        self.log = logger
        self.normalizer = RecordNormalizer(settings.get("extra_column_aliases") or {})
        self.excel_path = self._resolve_excel(settings.get("excel_file", ""))
        self.provider = ExcelDataProvider(self.excel_path, self.normalizer, settings.get("reader", "auto"), logger)
        self.search_engine = SearchEngine(self.provider, int(settings.get("max_tree_items", 150)))
        self.bom = BomBuilder(self.provider)
        mode = str(settings.get("default_mode", "ENGINEERING")).upper()
        self.mode = mode if mode in self.MODES else "ENGINEERING"
        self.node_map: dict[str, TreeNode] = {}
        self.current_result: SearchResult | None = None
        self.selected_record: Record | None = None   # full original Excel row of the selection
        self.selected_item: str | None = None
        self._loading = False
        self._queue: "queue.Queue" = queue.Queue()
        self._layout_cols = 0

        self.title(APP_TITLE)
        self.configure(bg=C_BG)
        self._init_geometry()
        self._init_fonts_and_styles()
        self._build_ui()
        self._bind_keys()
        self.report_callback_exception = self._on_tk_error
        self.protocol("WM_DELETE_WINDOW", self.quit_app)
        self.after(50, lambda: self.load_excel("startup"))

    # -- setup -----------------------------------------------------------------
    def _resolve_excel(self, configured: str) -> Path:
        path = resolve_path(configured) if configured else BASE_DIR / "data"
        if path.is_file():
            return path
        data_dir = BASE_DIR / "data"
        candidates = sorted(list(data_dir.glob("*.xlsm")) + list(data_dir.glob("*.xlsx"))) if data_dir.is_dir() else []
        candidates = [c for c in candidates if not c.name.startswith("~$")]
        if len(candidates) == 1:
            self.log.warning("Configured Excel not found (%s); using %s", configured, candidates[0].name)
            return candidates[0]
        return path

    def _init_geometry(self):
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        w = min(int(self.settings.get("window_width", 850)), sw - 40)
        h = min(int(self.settings.get("window_height", 850)), sh - 80)
        x, y = max(0, (sw - w) // 2), max(0, (sh - h) // 2 - 20)
        self.geometry(f"{w}x{h}+{x}+{y}")
        self.minsize(620, 500)

    def _init_fonts_and_styles(self):
        fam = pick_font(["Segoe UI", "Segoe UI Variable Text", "Helvetica Neue", "Arial", "DejaVu Sans"])
        mono = pick_font(["Consolas", "Cascadia Mono", "Lucida Console", "Courier New", "DejaVu Sans Mono"])
        self.fonts = {
            "title": (fam, 17, "bold"),
            "panel": (fam, 10, "bold"),
            "label": (fam, 9, "bold"),
            "value": (fam, 10),
            "value_b": (fam, 10, "bold"),
            "small": (fam, 8),
            "small_b": (fam, 8, "bold"),
            "entry": (fam, 11),
            "button": (fam, 10, "bold"),
            "tree": (fam, 10),
            "marking": (mono, 15, "bold"),
            "logo": (fam, 15, "bold"),
        }
        st = ttk.Style(self)
        if "clam" in st.theme_names():
            st.theme_use("clam")
        st.configure(".", font=self.fonts["value"], background=C_BG, foreground=C_TEXT)
        for name in ("Treeview", "Details.Treeview"):
            st.configure(name, background=C_PANEL, fieldbackground=C_PANEL, foreground=C_TEXT,
                         rowheight=24 if name == "Treeview" else 22, font=self.fonts["tree"],
                         bordercolor=C_PANEL, borderwidth=0)
            st.map(name, background=[("selected", C_SELECT)], foreground=[("selected", C_TEXT)])
        st.configure("Treeview.Heading", font=self.fonts["small_b"], background=C_SOFT, foreground=C_NAVY,
                     relief="flat", borderwidth=0)
        st.map("Treeview.Heading", background=[("active", "#CFD7DC")])
        st.configure("Tool.TButton", font=self.fonts["small_b"], padding=(10, 4), background="#F5F7F8",
                     foreground=C_NAVY, bordercolor=C_BORDER, relief="flat")
        st.map("Tool.TButton", background=[("active", C_SOFT), ("disabled", "#F0F0F0")])
        st.configure("TEntry", fieldbackground=C_PANEL, bordercolor=C_BORDER, lightcolor=C_BORDER, padding=4)
        st.configure("TCombobox", fieldbackground=C_PANEL, bordercolor=C_BORDER)
        st.configure("Vertical.TScrollbar", background=C_SOFT, troughcolor=C_BG, bordercolor=C_BG, arrowcolor=C_NAVY)
        st.configure("Horizontal.TScrollbar", background=C_SOFT, troughcolor=C_BG, bordercolor=C_BG, arrowcolor=C_NAVY)
        st.configure("TPanedwindow", background=C_BG)
        st.configure("Sash", sashthickness=6, gripcount=0)

    # -- UI ----------------------------------------------------------------------
    def _build_ui(self):
        f = self.fonts
        # Header -----------------------------------------------------------------
        header = tk.Frame(self, bg=C_PANEL, highlightthickness=1, highlightbackground=C_BORDER)
        header.pack(fill="x", padx=10, pady=(10, 6))
        title_row = tk.Frame(header, bg=C_PANEL)
        title_row.pack(fill="x", padx=14, pady=(10, 2))
        tk.Label(title_row, text=APP_TITLE, bg=C_PANEL, fg=C_NAVY, font=f["title"]).pack(side="left")
        logo = tk.Frame(title_row, bg=C_PANEL)
        logo.pack(side="right")
        tk.Label(logo, text="GP", bg=C_PANEL, fg=C_NAVY, font=f["logo"]).pack(side="left")
        tk.Label(logo, text="V", bg=C_PANEL, fg=C_GREEN, font=f["logo"]).pack(side="left")
        tk.Frame(header, bg=C_GREEN, height=2).pack(fill="x", padx=14, pady=(2, 6))

        search_row = tk.Frame(header, bg=C_PANEL)
        search_row.pack(fill="x", padx=14, pady=(0, 4))
        tk.Label(search_row, text="Assembly (Top level):", bg=C_PANEL, fg=C_TEXT, font=f["value"]).pack(side="left")
        self.search_var = tk.StringVar()
        self.entry = tk.Entry(search_row, textvariable=self.search_var, font=f["entry"], relief="solid", bd=1,
                              highlightthickness=1, highlightcolor=C_BLUE, highlightbackground=C_BORDER,
                              fg=C_TEXT, insertbackground=C_TEXT)
        self.entry.pack(side="left", fill="x", expand=True, padx=(8, 8), ipady=4)
        self.search_btn = tk.Button(search_row, text="SEARCH", command=self.do_search, bg=C_BLUE, fg="white",
                                    activebackground=C_BLUE_DARK, activeforeground="white", relief="flat",
                                    font=f["button"], padx=18, pady=4, cursor="hand2", bd=0)
        self.search_btn.pack(side="left")
        self.status_lbl = tk.Label(header, text="Item number / Assembly / Part number / MPN  -  Enter to search",
                                   bg=C_PANEL, fg=C_MUTED, font=f["small"], anchor="w")
        self.status_lbl.pack(fill="x", padx=14, pady=(0, 8))

        # Footer (packed before the body so it always stays visible) -------------
        footer = tk.Frame(self, bg=C_BG)
        footer.pack(side="bottom", fill="x", padx=12, pady=(0, 6))
        self.footer_lbl = tk.Label(footer, text="Loading Excel...", bg=C_BG, fg=C_MUTED, font=f["small"],
                                   anchor="w", justify="left")
        self.footer_lbl.pack(side="left", fill="x", expand=True)
        brand = tk.Frame(footer, bg=C_BG)
        brand.pack(side="right")
        tk.Label(brand, text="GP", bg=C_BG, fg="#9AA4AC", font=(f["logo"][0], 12, "bold")).pack(side="left")
        tk.Label(brand, text="V", bg=C_BG, fg=C_GREEN, font=(f["logo"][0], 12, "bold")).pack(side="left")
        self.footer_lbl.bind("<Configure>", lambda e: self.footer_lbl.configure(wraplength=max(200, e.width)))

        # Body: BOM (top) / information (bottom, scrollable) --------------------------
        self.paned = ttk.Panedwindow(self, orient="vertical")
        self.paned.pack(fill="both", expand=True, padx=10, pady=(0, 6))
        self.paned.add(self._build_bom_panel(self.paned), weight=3)
        self.paned.add(self._build_info_area(self.paned), weight=4)
        self.after(120, self._init_sash)
        self.entry.focus_set()

    def _init_sash(self):
        self.update_idletasks()
        total = self.paned.winfo_height()
        if total > 100:
            self.paned.sashpos(0, int(total * 0.46))

    def _build_bom_panel(self, master):
        f = self.fonts
        panel = tk.Frame(master, bg=C_PANEL, highlightthickness=1, highlightbackground=C_BORDER)
        head = tk.Frame(panel, bg=C_PANEL)
        head.pack(fill="x", padx=10, pady=(8, 0))
        # MODE is packed first so it keeps its space when the window is narrow.
        modes = tk.Frame(head, bg=C_PANEL)
        modes.pack(side="right", anchor="n")
        tk.Label(head, text="BOM STRUCTURE", bg=C_PANEL, fg=C_NAVY,
                 font=(f["panel"][0], 13, "bold")).pack(side="left", anchor="nw")
        self.view_lbl = tk.Label(panel, text="", bg=C_PANEL, fg=C_AMBER, font=f["small_b"], anchor="w",
                                 justify="left")
        self.view_lbl.pack(fill="x", padx=10)
        self.view_lbl.bind("<Configure>", lambda e: self.view_lbl.configure(wraplength=max(200, e.width)))
        tk.Label(modes, text="MODE", bg=C_PANEL, fg=C_NAVY, font=f["label"]).grid(row=0, column=0, padx=(0, 8), sticky="n", pady=(5, 0))
        self.mode_widgets = {}
        for i, m in enumerate(self.MODES):
            btn = tk.Label(modes, text=m, font=f["small_b"], padx=14, pady=4, cursor="hand2", bd=0)
            btn.grid(row=0, column=i + 1, padx=3)
            cap = tk.Label(modes, text="", bg=C_PANEL, font=f["small"])
            cap.grid(row=1, column=i + 1)
            btn.bind("<Button-1>", lambda e, mm=m: self.set_mode(mm))
            self.mode_widgets[m] = (btn, cap)
        self._paint_mode()

        tree_box = tk.Frame(panel, bg=C_PANEL)
        tree_box.pack(fill="both", expand=True, padx=10, pady=(4, 4))
        tree_box.rowconfigure(0, weight=1)
        tree_box.columnconfigure(0, weight=1)
        self.tree = ttk.Treeview(tree_box, columns=("info",), show="tree headings", selectmode="browse")
        self.tree.heading("#0", text="Item / Alternative / Attribute", anchor="w")
        self.tree.heading("info", text="Status / Source", anchor="w")
        self.tree.column("#0", width=470, minwidth=260, stretch=True)
        self.tree.column("info", width=190, minwidth=120, stretch=False)
        vs = ttk.Scrollbar(tree_box, orient="vertical", command=self.tree.yview)
        hs = ttk.Scrollbar(tree_box, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vs.grid(row=0, column=1, sticky="ns")
        hs.grid(row=1, column=0, sticky="ew")
        self.tree.tag_configure("item", font=(f["tree"][0], 10, "bold"), foreground=C_NAVY)
        self.tree.tag_configure("search", font=(f["tree"][0], 10, "bold"), foreground=C_NAVY)
        self.tree.tag_configure("use", foreground=C_OK)
        self.tree.tag_configure("dontuse", foreground=C_RED)
        self.tree.tag_configure("obsolete", foreground=C_AMBER)
        self.tree.tag_configure("attr", foreground="#3A454F")
        self.tree.tag_configure("muted", foreground=C_MUTED)
        self.tree.tag_configure("match", background="#EEF7E6")
        self.tree.bind("<<TreeviewSelect>>", self._on_tree_select)
        self.tree.bind("<Double-1>", self._on_tree_double)
        self.tree.bind("<Return>", self._on_tree_double)

        tools = tk.Frame(panel, bg=C_PANEL)
        tools.pack(fill="x", padx=10, pady=(0, 8))
        ttk.Button(tools, text="EXPAND ALL", style="Tool.TButton", command=lambda: self._expand_all(True)).pack(side="left")
        ttk.Button(tools, text="COLLAPSE ALL", style="Tool.TButton",
                   command=lambda: self._expand_all(False)).pack(side="left", padx=(6, 0))
        self.reload_btn = ttk.Button(tools, text="RELOAD EXCEL (F5)", style="Tool.TButton",
                                     command=lambda: self.load_excel("button"))
        self.reload_btn.pack(side="right")
        ttk.Button(tools, text="VIEW DETAILS", style="Tool.TButton", command=self.open_details).pack(side="right", padx=6)
        return panel

    def _build_info_area(self, master):
        outer = tk.Frame(master, bg=C_BG)
        outer.rowconfigure(0, weight=1)
        outer.columnconfigure(0, weight=1)
        self.info_canvas = tk.Canvas(outer, bg=C_BG, highlightthickness=0, bd=0)
        vs = ttk.Scrollbar(outer, orient="vertical", command=self.info_canvas.yview)
        self.info_canvas.configure(yscrollcommand=vs.set)
        self.info_canvas.grid(row=0, column=0, sticky="nsew")
        vs.grid(row=0, column=1, sticky="ns")
        self.info_inner = tk.Frame(self.info_canvas, bg=C_BG)
        self._inner_id = self.info_canvas.create_window((0, 0), window=self.info_inner, anchor="nw")
        self.info_inner.bind("<Configure>", lambda e: self.info_canvas.configure(
            scrollregion=self.info_canvas.bbox("all")))
        self.info_canvas.bind("<Configure>", self._on_info_resize)
        for w in (self.info_canvas, self.info_inner):
            w.bind("<Enter>", lambda e: self._wheel(True))
            w.bind("<Leave>", lambda e: self._wheel(False))

        f = self.fonts
        self.p_assembly = InfoPanel(self.info_inner, "ASSEMBLY INFORMATION", f)
        self.p_component = InfoPanel(self.info_inner, "PCB / COMPONENT INFORMATION", f)
        self.p_marking = MarkingPanel(self.info_inner, f)
        self.p_material = InfoPanel(self.info_inner, "MATERIAL INFORMATION", f, columns=2)
        self.p_alts = InfoPanel(self.info_inner, "ALTERNATIVES / COMPONENT INFORMATION", f)
        self.p_extra = InfoPanel(self.info_inner, "ADDITIONAL FIELDS (EXCEL)", f)
        self._layout_panels(1)
        self._clear_panels()
        return outer

    def _on_info_resize(self, event):
        self.info_canvas.itemconfigure(self._inner_id, width=event.width)
        self._layout_panels(2 if event.width >= 980 else 1)

    def _layout_panels(self, cols):
        if cols == self._layout_cols:
            return
        self._layout_cols = cols
        panels = [self.p_assembly, self.p_component, self.p_marking, self.p_material, self.p_alts, self.p_extra]
        for p in panels:
            p.grid_forget()
        for c in range(2):
            self.info_inner.columnconfigure(c, weight=1 if c < cols else 0, uniform="info" if c < cols else "")
        visible = [p for p in panels if getattr(p, "_visible", True)]
        for i, p in enumerate(visible):
            r, c = divmod(i, cols)
            p.grid(row=r, column=c, sticky="nsew", padx=(0 if c == 0 else 3, 0 if c == cols - 1 else 3), pady=3)

    def _set_panel_visible(self, panel, visible: bool):
        if getattr(panel, "_visible", True) != visible:
            panel._visible = visible
            cols = self._layout_cols
            self._layout_cols = 0
            self._layout_panels(cols)

    def _wheel(self, on: bool):
        if on:
            self.bind_all("<MouseWheel>", self._on_wheel)
            self.bind_all("<Button-4>", self._on_wheel)
            self.bind_all("<Button-5>", self._on_wheel)
        else:
            self.unbind_all("<MouseWheel>")
            self.unbind_all("<Button-4>")
            self.unbind_all("<Button-5>")

    def _on_wheel(self, event):
        if getattr(event, "num", None) == 4:
            step = -2
        elif getattr(event, "num", None) == 5:
            step = 2
        else:
            step = -1 * int(event.delta / 120) if abs(event.delta) >= 120 else (-1 if event.delta > 0 else 1)
        if self.info_canvas.yview() != (0.0, 1.0):
            self.info_canvas.yview_scroll(step, "units")

    def _bind_keys(self):
        self.entry.bind("<Return>", lambda e: self.do_search())
        self.entry.bind("<KP_Enter>", lambda e: self.do_search())
        self.bind("<F5>", lambda e: self.load_excel("F5"))
        self.bind("<Escape>", lambda e: self.clear_search())
        self.bind("<Control-f>", lambda e: self.focus_search())
        self.bind("<Control-F>", lambda e: self.focus_search())
        self.bind("<Control-q>", lambda e: self.quit_app())
        self.bind("<Control-Q>", lambda e: self.quit_app())

    # -- mode ----------------------------------------------------------------------
    def _paint_mode(self):
        for m, (btn, cap) in self.mode_widgets.items():
            active = m == self.mode
            btn.configure(bg=C_NAVY if active else C_SOFT, fg="white" if active else C_MUTED)
            cap.configure(text="active" if active else "inactive", fg=C_NAVY if active else "#A0AAB2")

    def set_mode(self, mode: str):
        if mode == self.mode:
            return
        self.mode = mode
        self._paint_mode()
        self.log.info("Mode changed to %s", mode)
        if self.current_result is not None and self.current_result.found:
            self._render_result(self.current_result, keep_selection=True)
        else:
            self._clear_panels()
        self.focus_search()

    # -- loading -------------------------------------------------------------------
    def load_excel(self, reason: str = ""):
        if self._loading:
            return
        self._loading = True
        self.reload_btn.state(["disabled"])
        self.search_btn.configure(state="disabled")
        self.status_lbl.configure(text="Reading Excel from disk...", fg=C_BLUE)
        self.footer_lbl.configure(text=f"Loading {self.excel_path.name} ...")
        self.configure(cursor="watch")
        self.log.info("Loading Excel (%s): %s", reason, self.excel_path)

        def worker():
            try:
                fresh = ExcelDataProvider(self.excel_path, self.normalizer,
                                          self.settings.get("reader", "auto"), self.log).load_workbook()
                self._queue.put(("ok", fresh))
            except ExcelLoadError as exc:
                self.log.error("Load failed: %s", str(exc).replace("\n", " "))
                self._queue.put(("error", str(exc)))
            except Exception as exc:  # unexpected -> log traceback, short message to operator
                self.log.exception("Unexpected error while reading the Excel")
                self._queue.put(("error", f"The Excel file could not be read ({exc.__class__.__name__}).\n"
                                          "Details were written to logs/app.log"))

        threading.Thread(target=worker, daemon=True).start()
        self.after(100, self._poll_load)

    def _poll_load(self):
        try:
            status, payload = self._queue.get_nowait()
        except queue.Empty:
            self.after(100, self._poll_load)
            return
        self._loading = False
        self.configure(cursor="")
        self.reload_btn.state(["!disabled"])
        self.search_btn.configure(state="normal")
        if status == "ok":
            self.provider = payload
            self.search_engine.provider = payload
            self.bom.provider = payload
            self._update_footer()
            self.view_lbl.configure(
                text=self.bom.view_name + ("  -  Item → Alternatives (not a FactoryLogix genealogy)"
                                           if self.provider.hierarchy_mode == "logical" else ""),
                fg=C_AMBER if self.provider.hierarchy_mode == "logical" else C_OK)
            if self.search_var.get().strip():
                self.do_search(log_reason="reload")
            else:
                self._show_welcome()
                self.status_lbl.configure(
                    text=f"Excel loaded: {len(self.provider.records):,} rows, "
                         f"{self.provider.item_count():,} item numbers. Scan or type a number and press Enter.",
                    fg=C_OK)
        else:
            self.footer_lbl.configure(text=f"Source: {self.excel_path.name}  ·  NOT LOADED  ·  "
                                           f"{dt.datetime.now():%H:%M:%S}")
            self.status_lbl.configure(text=payload.replace("\n", " "), fg=C_RED)
            messagebox.showerror(APP_TITLE, payload, parent=self)
        self.focus_search()

    def _update_footer(self):
        p = self.provider
        mtime = f"{p.file_mtime:%Y-%m-%d %H:%M}" if p.file_mtime else EMPTY
        loaded = f"{p.loaded_at:%H:%M:%S}" if p.loaded_at else EMPTY
        self.footer_lbl.configure(
            text=f"Source: {self.excel_path.name}   ·   Sheet: {p.sheet_summary()}   ·   "
                 f"Rows loaded: {len(p.records):,}   ·   Last reload: {loaded}   ·   "
                 f"File saved: {mtime}   ·   Reader: {p.reader_used}   ·   Read-only")

    # -- search --------------------------------------------------------------------
    def focus_search(self):
        self.entry.focus_set()
        self.entry.select_range(0, "end")
        self.entry.icursor("end")

    def clear_search(self):
        self.search_var.set("")
        self.current_result = None
        self._show_welcome()
        self.status_lbl.configure(text="Search cleared.", fg=C_MUTED)
        self.focus_search()

    def do_search(self, log_reason: str = "search"):
        if self._loading:
            # The text stays in the box; _poll_load runs the search when loading ends.
            self.status_lbl.configure(text="Excel is loading - the search will run automatically.", fg=C_BLUE)
            self.focus_search()
            return
        raw = self.search_var.get()
        query = "".join(ch for ch in raw if ch.isprintable()).strip()
        if query != raw:
            self.search_var.set(query)
        if not query:
            self.status_lbl.configure(text="Type or scan an Item number / Assembly / Part number / MPN.", fg=C_MUTED)
            self.focus_search()
            return
        if not self.provider.records:
            self.status_lbl.configure(text="No Excel data loaded. Press F5 to reload.", fg=C_RED)
            self.focus_search()
            return
        res = self.search_engine.search(query)
        self.current_result = res
        n_items = len(res.items)
        n_rows = sum(len(i["matched"]) for i in res.items.values())
        self.log.info("Search (%s) '%s' mode=%s -> %s match, %d item(s), %d row(s), %.1f ms",
                      log_reason, query, self.mode, res.match_kind, n_items, n_rows, res.elapsed_ms)
        if not res.found:
            self._show_message_tree(f"No match for '{query}'")
            self._clear_panels()
            self.status_lbl.configure(text=f"'{query}' was not found in Item number / MPN.", fg=C_RED)
            self.bell()
        else:
            self._render_result(res)
            kind = {"exact": "Exact match", "prefix": "Starts with", "contains": "Contains"}[res.match_kind]
            fields = sorted({SearchEngine.FIELD_LABELS.get(i["field"], i["field"]) for i in res.items.values()})
            more = f" (+{res.truncated} not listed)" if res.truncated else ""
            self.status_lbl.configure(
                text=f"{kind}: {n_items} item(s), {n_rows} row(s) by {', '.join(fields)}{more}   "
                     f"[{res.elapsed_ms:.0f} ms]",
                fg=C_OK if res.match_kind == "exact" else C_AMBER)
        self.focus_search()

    # -- tree ------------------------------------------------------------------------
    def _render_result(self, res: SearchResult, keep_selection=False):
        previous = self.tree.selection()[0] if keep_selection and self.tree.selection() else None
        prev_uid = None
        if previous and previous in self.node_map:
            n = self.node_map[previous]
            prev_uid = n.record.uid if n.record else ("item:" + (n.item_key or ""))
        root = self.bom.build(res, self.mode)
        self._fill_tree(root)
        target = None
        if prev_uid:
            for iid, n in self.node_map.items():
                uid = n.record.uid if n.record else ("item:" + (n.item_key or ""))
                if uid == prev_uid:
                    target = iid
                    break
        if target is None:
            matched = res.matched_uids()
            target = next((iid for iid, n in self.node_map.items()
                           if n.kind == "alt" and n.record is not None and n.record.uid in matched), None)
        if target is None:
            target = self._first_iid_of(root)
        if target:
            self.tree.selection_set(target)
            self.tree.focus(target)
            self.tree.see(target)

    def _first_iid_of(self, root: TreeNode):
        for iid, n in self.node_map.items():
            if n is root:
                if root.kind == "search" and root.children and root.children[0].kind == "item":
                    return next((i for i, m in self.node_map.items() if m is root.children[0]), iid)
                return iid
        return None

    def _fill_tree(self, root: TreeNode):
        self.tree.delete(*self.tree.get_children())
        self.node_map = {}
        counter = [0]

        def add(parent_iid, node):
            counter[0] += 1
            iid = f"n{counter[0]}"
            tags = list(node.tags) + ([node.kind] if node.kind in ("item", "search") else [])
            self.tree.insert(parent_iid, "end", iid=iid, text=node.text, values=(node.info,),
                             open=node.open, tags=tags)
            self.node_map[iid] = node
            for ch in node.children:
                add(iid, ch)
        add("", root)

    def _show_message_tree(self, text):
        self.tree.delete(*self.tree.get_children())
        self.node_map = {}
        self.tree.insert("", "end", iid="msg", text=text, values=("",), tags=("muted",))

    def _show_welcome(self):
        self._show_message_tree("Scan or type an Item number / Assembly / MPN and press Enter")
        self._clear_panels()

    def _expand_all(self, expand: bool):
        def walk(iid):
            self.tree.item(iid, open=expand)
            for ch in self.tree.get_children(iid):
                walk(ch)
        for top in self.tree.get_children(""):
            walk(top)
            if not expand:
                self.tree.item(top, open=True)

    def _on_tree_select(self, _event=None):
        sel = self.tree.selection()
        node = self.node_map.get(sel[0]) if sel else None
        if node is None:
            return
        if node.record is not None:
            self.show_record(node.record)
        elif node.item_key:
            self.show_item(node.item_key)
        elif node.kind == "search" and self.current_result:
            self.show_search_summary(self.current_result)

    def _on_tree_double(self, event=None):
        sel = self.tree.selection()
        node = self.node_map.get(sel[0]) if sel else None
        if node is None:
            return None
        if node.kind in ("alt", "attr", "ref") and node.record is not None:
            self.open_details(node.record.get("item_number") or "")
            return "break"
        return None  # items / groups: default expand-collapse behaviour

    # -- information panels ------------------------------------------------------------
    def _clear_panels(self):
        self.selected_record = None
        self.selected_item = None
        self.p_assembly.set_rows([])
        self.p_component.set_rows([])
        self.p_marking.set_markings([], "")
        self.p_material.set_rows([])
        self.p_alts.set_rows([])
        self._set_panel_visible(self.p_alts, False)
        self._set_panel_visible(self.p_extra, False)
        self._set_panel_visible(self.p_material, self.mode == "ENGINEERING")

    def _field_label(self, rec_or_table_headers: dict, key: str) -> str:
        label = RecordNormalizer.label_for(key)
        original = rec_or_table_headers.get(key)
        if original and RecordNormalizer.normalize_header(original) != RecordNormalizer.normalize_header(label):
            return f"{label} ({original})"
        return label

    def _rec_rows(self, rec: Record, keys, styles=None):
        rows = []
        for key in keys:
            if key in rec.headers_by_key:  # show only columns that exist in that sheet
                style = (styles or {}).get(key)
                rows.append((self._field_label(rec.headers_by_key, key), rec.get(key), style))
        return rows

    @staticmethod
    def _use_text(rec: Record):
        flag = rec.status_flag()
        raw = rec.get("status")
        if flag == "YES":
            return f"{raw}  →  DON'T USE / NO USAR", "bad"
        if flag == "NO":
            return f"{raw}  →  USE / USAR", "good"
        return raw, None

    def show_record(self, rec: Record):
        """Selection of an alternative / attribute / row: the whole original row is kept."""
        self.selected_record = rec
        self.selected_item = rec.get("item_number")
        prod = self.mode == "PRODUCTION"
        # ASSEMBLY INFORMATION
        rows = self._rec_rows(rec, ["item_number", "text"], {"item_number": "strong"})
        if "status" in rec.headers_by_key:
            txt, style = self._use_text(rec)
            rows.append((self._field_label(rec.headers_by_key, "status"), txt, style))
        obs_style = "warn" if RecordNormalizer.flag(rec.get("obsolete")) == "YES" else None
        rows += self._rec_rows(rec, ["obsolete"], {"obsolete": obs_style})
        if not prod:
            rows += self._rec_rows(rec, ["valid_until", "temperature", "duration", "parent", "level", "quantity"])
            rows.append(("SOURCE", f"Sheet '{rec.sheet}', row {rec.row}", "muted"))
        self.p_assembly.set_rows(rows)
        # PCB / COMPONENT INFORMATION
        comp = self._rec_rows(rec, ["item_number", "alternative", "manufacturer", "mpn", "marking"],
                              {"mpn": "strong"})
        if not rec.has_detail:
            comp.append(("NOTE", f"Sheet '{rec.sheet}' row {rec.row} contains only the Item number; "
                                 "all other columns are empty in the Excel.", "warn"))
        self.p_component.set_rows(comp)
        # MARKING
        if "marking" in rec.headers_by_key:
            alt = rec.get("alternative")
            prefix = " · ".join(x for x in (f"ALT {alt}" if alt else "", rec.get("manufacturer") or "",
                                                  rec.get("mpn") or "") if x)
            self.p_marking.set_markings([(prefix, rec.get("marking"),
                                          "bad" if rec.status_flag() == "YES" else None)],
                                        f"{rec.sheet}!{rec.row}")
        else:
            self.p_marking.set_markings([], f"{rec.sheet}!{rec.row}")
        # MATERIAL INFORMATION (engineering)
        self.p_material.set_rows(self._rec_rows(
            rec, ["lead_free", "received_lead_free", "temperature", "valid_until", "duration"]))
        self._set_panel_visible(self.p_material, not prod)
        # other alternatives of the same item
        self._fill_alternatives(rec.get("item_number"), highlight=rec)
        # ADDITIONAL FIELDS: real columns without a known concept
        extra = [(h, v, None) for h, v in rec.values.items() if h not in rec.headers_by_key.values()]
        self.p_extra.set_rows(extra)
        self._set_panel_visible(self.p_extra, bool(extra) and not prod)

    def _fill_alternatives(self, item, highlight: Record | None = None):
        if not item:
            self._set_panel_visible(self.p_alts, False)
            return
        recs = [r for r in self.provider.get_alternatives(item) if r.has_detail]
        if self.mode == "PRODUCTION":
            recs = [r for r in recs if r.status_flag() != "YES"]
        rows = []
        for r in recs:
            alt = r.get("alternative") or f"row {r.row}"
            desc = "  ·  ".join(RecordNormalizer.display(r.get(k)) for k in ("manufacturer", "mpn")
                                    if k in r.headers_by_key)
            if "marking" in r.headers_by_key:
                desc += f"\nMarking: {RecordNormalizer.display(r.get('marking'))}"
            flag = r.status_flag()
            if flag in ("YES", "NO"):
                desc += "   [" + ("DON'T USE" if flag == "YES" else "USE") + "]"
            style = "bad" if flag == "YES" else ("strong" if highlight is not None and r is highlight else None)
            rows.append((f"ALT {alt}" + (" ◀" if r is highlight else ""), desc, style))
        refs = [r for r in self.provider.get_alternatives(item) if not r.has_detail]
        if refs and self.mode == "ENGINEERING":
            sheets = OrderedDict()
            for r in refs:
                sheets.setdefault(r.sheet, []).append(str(r.row))
            for s, rws in sheets.items():
                rows.append((f"'{s}' ROWS", f"{len(rws)} row(s) with Item number only: {', '.join(rws[:30])}"
                             + (" ..." if len(rws) > 30 else ""), "muted"))
        self.p_alts.set_rows(rows)
        self._set_panel_visible(self.p_alts, bool(rows))

    def show_item(self, item_key: str):
        recs_all = self.provider._item_index.get(item_key, [])
        if not recs_all:
            return
        item = recs_all[0].get("item_number")
        self.selected_item = item
        self.selected_record = None
        recs = self.provider.get_alternatives(item)
        details = [r for r in recs if r.has_detail]
        refs = [r for r in recs if not r.has_detail]
        prod = self.mode == "PRODUCTION"
        visible = [r for r in details if not (prod and r.status_flag() == "YES")]
        headers = {}
        for r in recs:
            for k, h in r.headers_by_key.items():
                headers.setdefault(k, h)

        def summary(key):
            vals = OrderedDict((r.get(key), None) for r in details if r.get(key) is not None)
            if not vals:
                return None
            if len(vals) == 1:
                return next(iter(vals))
            return "Varies by alternative: " + " | ".join(list(vals)[:8])

        # Validation exactly like the macro (G2): any Status/Stopped = "Yes" -> O (NO USAR)
        if not details:
            validation, vstyle = "No detail data in Excel for this item", "warn"
        elif any(r.status_flag() == "YES" for r in details):
            validation, vstyle = "✘ O  -  NO USAR (an alternative has Status = Yes)", "bad"
        else:
            validation, vstyle = "✔ P  -  USAR (no alternative with Status = Yes)", "good"
        rows = [(self._field_label(headers, "item_number"), item, "strong")]
        if "text" in headers:
            rows.append((self._field_label(headers, "text"), summary("text"), None))
        rows.append(("SMT VALIDATION", validation, vstyle))
        if "obsolete" in headers:
            obs = [r for r in details if RecordNormalizer.flag(r.get("obsolete")) == "YES"]
            rows.append((self._field_label(headers, "obsolete"),
                         (f"Yes in {len(obs)} of {len(details)} alternative(s)" if obs else summary("obsolete")),
                         "warn" if obs else None))
        if not prod:
            for key in ("valid_until", "temperature", "duration"):
                if key in headers:
                    rows.append((self._field_label(headers, key), summary(key), None))
            src = OrderedDict()
            for r in recs:
                src.setdefault(r.sheet, 0)
                src[r.sheet] += 1
            rows.append(("SOURCE", ", ".join(f"'{s}' ({n} row(s))" for s, n in src.items()), "muted"))
        self.p_assembly.set_rows(rows)

        comp = [(self._field_label(headers, "item_number"), item, "strong"),
                ("ALTERNATIVES", f"{len(details)} with data" + (f", {len(refs)} row(s) without detail" if refs else ""),
                 None)]
        for key in ("manufacturer", "mpn"):
            if key in headers:
                vals = list(OrderedDict((r.get(key), None) for r in visible if r.get(key)))
                comp.append((self._field_label(headers, key), ", ".join(vals) if vals else None, None))
        if not details and refs:
            comp.append(("NOTE", "The Excel only contains the Item number for this part "
                                 "(Alternative, Manufacturer, MPN, Marking... are empty).", "warn"))
        self.p_component.set_rows(comp)

        marks = []
        for r in visible:
            if "marking" not in r.headers_by_key:
                continue
            prefix = " · ".join(x for x in (f"ALT {r.get('alternative')}" if r.get("alternative") else f"ROW {r.row}",
                                                  r.get("manufacturer") or "", r.get("mpn") or "") if x)
            marks.append((prefix, r.get("marking"), "bad" if r.status_flag() == "YES" else None))
        self.p_marking.set_markings(marks, f"{len(marks)} alternative(s)" if marks else "")

        self.p_material.set_rows([(self._field_label(headers, k), summary(k), None) for k in
                                  ("lead_free", "received_lead_free", "temperature", "valid_until", "duration")
                                  if k in headers])
        self._set_panel_visible(self.p_material, not prod)
        self._fill_alternatives(item)
        self._set_panel_visible(self.p_extra, False)

    def show_search_summary(self, res: SearchResult):
        self.selected_record = None
        rows = [("SEARCH", res.query, "strong"),
                ("MATCH", {"exact": "Exact", "prefix": "Starts with", "contains": "Contains"}.get(res.match_kind, EMPTY), None),
                ("ITEMS", str(len(res.items)), None)]
        self.p_assembly.set_rows(rows)
        self.p_component.set_rows([(f"#{i + 1}", info["item"], None) for i, info in enumerate(list(res.items.values())[:25])])
        self.p_marking.set_markings([], "")
        self.p_material.set_rows([])
        self._set_panel_visible(self.p_alts, False)
        self._set_panel_visible(self.p_extra, False)

    # -- misc --------------------------------------------------------------------------
    def open_details(self, item_filter: str | None = None):
        if not self.provider.records:
            messagebox.showinfo(APP_TITLE, "No Excel data loaded.", parent=self)
            return
        if item_filter is None:
            item_filter = self.selected_item or ""
            if not item_filter and self.current_result and len(self.current_result.items) == 1:
                item_filter = next(iter(self.current_result.items.values()))["item"]
        DetailsWindow(self, item_filter)

    def _on_tk_error(self, exc, val, tb):
        self.log.error("UI error", exc_info=(exc, val, tb))
        if self.winfo_exists():
            self.status_lbl.configure(text="An internal error occurred. Details in logs/app.log", fg=C_RED)

    def quit_app(self):
        self.log.info("Application closed")
        self.destroy()


# ============================================================================
# Console check (python GPV_Product_Configuration.py --check)
# ============================================================================

def run_check(settings: dict, logger: logging.Logger, reader: str | None = None, queries=None) -> int:
    nz = RecordNormalizer(settings.get("extra_column_aliases") or {})
    path = resolve_path(settings.get("excel_file", ""))
    prov = ExcelDataProvider(path, nz, reader or settings.get("reader", "auto"), logger)
    try:
        prov.load_workbook()
    except ExcelLoadError as exc:
        print("ERROR:", exc)
        return 1
    print(f"File     : {path.name}")
    print(f"Reader   : {prov.reader_used}   ({prov.load_seconds:.1f}s)")
    print(f"Sheets   : {prov.list_sheets()}")
    for t in prov.tables.values():
        print(f"\n[{t.name}] scanned rows={t.scanned_rows:,} header row={t.header_row} records={len(t.records):,}")
        for idx, h in t.columns:
            key = t.key_by_header.get(h, "-")
            filled = sum(1 for r in t.records if r.values.get(h) is not None)
            print(f"   col {idx + 1:>2} {h!r:<24} -> {key:<20} filled={filled:,}")
        if t.ignored_columns:
            print(f"   ignored (outside data block): {t.ignored_columns}")
        for r, vals in t.preamble:
            print(f"   preamble row {r}: {vals}")
    print(f"\nRecords  : {len(prov.records):,}   distinct item numbers: {prov.item_count():,}")
    print(f"Hierarchy: {prov.hierarchy_mode}")
    eng = SearchEngine(prov)
    for q in queries or []:
        res = eng.search(q)
        print(f"\nSEARCH {q!r}: {res.match_kind}, items={len(res.items)}")
        for info in list(res.items.values())[:5]:
            for r in prov.get_alternatives(info["item"]):
                print("   ", r.uid, {k: v for k, v in r.fields.items() if v is not None})
    return 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=APP_TITLE)
    parser.add_argument("--check", action="store_true", help="print a data report and exit")
    parser.add_argument("--reader", choices=["auto", "openpyxl", "builtin"], help="override reader")
    parser.add_argument("queries", nargs="*", help="queries to test with --check")
    args = parser.parse_args(argv)
    logger = setup_logging()
    settings = load_settings()
    if args.reader:
        settings["reader"] = args.reader
    if args.check:
        return run_check(settings, logger, args.reader, args.queries)
    if tk is None:
        print("tkinter is not available in this Python installation.\n"
              "Install Python from python.org (tkinter is included by default).")
        return 1
    logger.info("Application start v%s | Python %s | openpyxl=%s", APP_VERSION, sys.version.split()[0],
                getattr(openpyxl, "__version__", "not installed"))
    try:
        app = GPVProductConfigurationApp(settings, logger)
    except tk.TclError as exc:
        logger.error("Could not start the window: %s", exc)
        print("The window could not be opened (no display available).")
        return 1
    app.mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())
